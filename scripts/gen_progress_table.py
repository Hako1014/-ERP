# -*- coding: utf-8 -*-
"""生成《电商 ERP 系统 · 项目开发进度表》。

以「项目开发进度表-通用模板.xlsx」为样式基底，逐任务填入 E:\\ERP 项目进度：
- 保留模板的 4 张表（使用说明 / 开发进度表 / 进度统计 / 状态说明）
- 保留表头配色、下拉校验、条件格式、统计公式等全部版式
- 开发进度表：26 行示例 → 74 行（TASK-001 ~ TASK-074）
- 进度统计：6 个阶段 → 12 个阶段

用法：python gen_progress_table.py [模板路径] [输出路径]
"""
import copy
import re
import sys
import zipfile
from xml.sax.saxutils import escape

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

SRC = sys.argv[1] if len(sys.argv) > 1 else r"C:\Users\13244\Desktop\项目开发进度表-通用模板.xlsx"
DST = sys.argv[2] if len(sys.argv) > 2 else r"E:\ERP\docs\项目开发进度表.xlsx"

PROJECT_TITLE = "电商 ERP 系统 · 项目开发进度表"

# 阶段名 -> 浅色底（逐阶段区分，沿用模板的浅色系风格）
STAGE_FILL = {
    "阶段 0 · 项目初始化": "EDE9FE",
    "阶段 1 · 基础架构": "DBEAFE",
    "阶段 2 · 数据库和数据模型": "CFFAFE",
    "阶段 3 · 后端基础能力": "D1FAE5",
    "阶段 4 · 核心业务模块": "FEF3C7",
    "阶段 5 · 前端基础能力": "FCE7F3",
    "阶段 6 · 业务页面": "E0F2FE",
    "阶段 7 · 权限和安全": "FEE2E2",
    "阶段 8 · 第三方服务": "FEF9C3",
    "阶段 9 · 测试": "F3E8FF",
    "阶段 10 · 部署": "DCFCE7",
    "阶段 11 · 优化和维护": "E5E7EB",
}

# (阶段, 任务卡范围, [(编号, 任务名称), ...])
STAGES = [
    ("阶段 0 · 项目初始化", "TASK-001–002", [
        (1, "项目初始化与工程脚手架"),
        (2, "开发规范与横切基础能力"),
    ]),
    ("阶段 1 · 基础架构", "TASK-003–008", [
        (3, "数据库会话与迁移框架"),
        (4, "审计日志与操作留痕"),
        (5, "单据编号生成服务"),
        (6, "文件上传服务"),
        (7, "Excel 导入导出服务"),
        (8, "PDF 生成与打印服务"),
    ]),
    ("阶段 2 · 数据库和数据模型", "TASK-009–016", [
        (9, "组织与权限数据模型"),
        (10, "商品与基础资料数据模型"),
        (11, "往来单位数据模型"),
        (12, "仓库与库位数据模型"),
        (13, "采购业务数据模型"),
        (14, "库存业务数据模型"),
        (15, "订单、履约与平台物流数据模型"),
        (16, "财务与审批数据模型"),
    ]),
    ("阶段 3 · 后端基础能力", "TASK-017–019", [
        (17, "登录认证与 JWT"),
        (18, "RBAC 权限框架"),
        (19, "通用审批流引擎"),
    ]),
    ("阶段 4 · 核心业务模块", "TASK-020–047", [
        (20, "商品资料与商品设置服务"),
        (21, "供应商与客户服务"),
        (22, "公司设置服务"),
        (23, "店铺对接与物流工具服务"),
        (24, "采购计划服务"),
        (25, "采购订单服务"),
        (26, "采购预付款申请服务"),
        (27, "采购退货单服务"),
        (28, "采购发票服务"),
        (29, "仓库与库位服务"),
        (30, "采购入库与商品上架服务"),
        (31, "采购出库服务"),
        (32, "调拨服务"),
        (33, "库存查询与库位调整服务"),
        (34, "库存盘点服务"),
        (35, "店铺库存与库存结存服务"),
        (36, "订单抓取与审核服务"),
        (37, "波次管理服务"),
        (38, "打印服务"),
        (39, "扫描发货服务"),
        (40, "订单售后与退货管理服务"),
        (41, "审批管理与我的审批服务"),
        (42, "会计科目与项目核算服务"),
        (43, "凭证管理服务"),
        (44, "付款单与收款单服务"),
        (45, "费用单/结算单/收入单服务"),
        (46, "发票管理与供应商账户余额服务"),
        (47, "报表与看板聚合服务"),
    ]),
    ("阶段 5 · 前端基础能力", "TASK-048–052", [
        (48, "前端工程与布局壳"),
        (49, "通用组件库"),
        (50, "登录页与路由守卫"),
        (51, "通用业务组件"),
        (52, "图表与看板基础"),
    ]),
    ("阶段 6 · 业务页面", "TASK-053–060", [
        (53, "基础设置页面"),
        (54, "商品管理页面"),
        (55, "采购管理页面"),
        (56, "库存管理页面"),
        (57, "订单管理页面"),
        (58, "协同办公页面"),
        (59, "财务管理页面"),
        (60, "首页看板页面"),
    ]),
    ("阶段 7 · 权限和安全", "TASK-061–063", [
        (61, "数据权限与 IP 访问控制"),
        (62, "接口鉴权与越权安全测试"),
        (63, "安全加固与权限回归"),
    ]),
    ("阶段 8 · 第三方服务", "TASK-064–067", [
        (64, "金蝶财务系统对接"),
        (65, "平台订单抓取对接"),
        (66, "物流对接"),
        (67, "短信与第三方仓储对接"),
    ]),
    ("阶段 9 · 测试", "TASK-068–070", [
        (68, "后端测试补齐"),
        (69, "前端与 E2E 测试"),
        (70, "关键业务链路联调"),
    ]),
    ("阶段 10 · 部署", "TASK-071–072", [
        (71, "部署与 CI/CD"),
        (72, "生产环境配置与发布"),
    ]),
    ("阶段 11 · 优化和维护", "TASK-073–074", [
        (73, "性能优化"),
        (74, "文档与维护"),
    ]),
]

# 仅「开发中」的任务（其余 DONE → 已测试）
IN_PROGRESS = {72}

# 例外说明（来自 docs/遗留项与二期清单.md 与各任务卡状态行）
NOTES = {
    64: "二期：仅金蝶适配接口 + 本地 mock",
    65: "二期：仅统一适配接口 + mock 抓单",
    66: "二期：仅物流接口 + mock 运单号",
    67: "二期：仅接口 + mock，未接真实短信/仓储",
    71: "CI 已配置，待 GitHub Actions 实跑",
    72: "未真实发布：缺域名/证书/服务器",
    73: "容量验收待补：本地仅 6 条订单",
    74: "OpenAPI 3.1：340 路径 / 493 操作",
}

STATUSES = ["未开发", "开发中", "已开发", "已测试"]

# 展开为行
rows = []
for stage, _rng, tasks in STAGES:
    for num, name in tasks:
        status = "开发中" if num in IN_PROGRESS else "已测试"
        rows.append((stage, f"TASK-{num:03d}", name, status, NOTES.get(num, "")))

assert len(rows) == 74, len(rows)

wb = load_workbook(SRC)

# ---------------------------------------------------------------- 使用说明
ws = wb["使用说明"]
ws["A1"] = PROJECT_TITLE

# ------------------------------------------------------------ 开发进度表
ws = wb["开发进度表"]
# 以第 2 行各列为样式模板
tpl_style = {col: copy.copy(ws[f"{col}2"]._style) for col in "ABCDE"}

for r, (stage, card, name, status, note) in enumerate(rows, start=2):
    vals = {"A": stage, "B": card, "C": name, "D": status, "E": note}
    for col, val in vals.items():
        c = ws[f"{col}{r}"]
        c._style = copy.copy(tpl_style[col])
        c.value = val
    # 阶段列按阶段着色
    ws[f"A{r}"].fill = PatternFill(fill_type="solid",
                                   start_color=STAGE_FILL[stage],
                                   end_color=STAGE_FILL[stage])
    ws.row_dimensions[r].height = 26

# 清掉模板遗留的多余行（模板示例到 27 行，新数据到 75 行，无需清）
ws.column_dimensions["A"].width = 24  # 本项目阶段名较长
ws.column_dimensions["B"].width = 11  # TASK-001
ws.column_dimensions["C"].width = 44
ws.column_dimensions["D"].width = 11
ws.column_dimensions["E"].width = 34
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:E{len(rows) + 1}"

# -------------------------------------------------------------- 进度统计
ws = wb["进度统计"]
# 整体进度：公式沿用模板（引用 开发进度表!D:D / B:B）
for i, status in enumerate(STATUSES):
    r = 4 + i
    ws[f"A{r}"] = status
    ws[f"B{r}"] = f"=COUNTIF(开发进度表!D:D,A{r})"
    ws[f"C{r}"] = f'=TEXT(COUNTIF(开发进度表!D:D,A{r})/MAX(1,COUNTA(开发进度表!B:B)-1),"0.0%")'
ws["A8"] = "合计"
ws["B8"] = "=SUM(B4:B7)"

# 阶段进度：模板为 6 行（12–17），本项目 12 行（12–23）
tpl_cells = {col: copy.copy(ws[f"{col}12"]._style) for col in "ABCDEFGH"}
for i, (stage, rng, tasks) in enumerate(STAGES):
    r = 12 + i
    ws[f"A{r}"] = stage
    ws[f"B{r}"] = rng
    ws[f"C{r}"] = f"=COUNTIF(开发进度表!A:A,A{r})"
    ws[f"D{r}"] = f'=COUNTIFS(开发进度表!A:A,A{r},开发进度表!D:D,"已测试")'
    ws[f"E{r}"] = f'=COUNTIFS(开发进度表!A:A,A{r},开发进度表!D:D,"已开发")'
    ws[f"F{r}"] = f'=COUNTIFS(开发进度表!A:A,A{r},开发进度表!D:D,"开发中")'
    ws[f"G{r}"] = f'=COUNTIFS(开发进度表!A:A,A{r},开发进度表!D:D,"未开发")'
    ws[f"H{r}"] = f'=TEXT(D{r}/MAX(1,C{r}),"0.0%")'
    for col in "ABCDEFGH":
        c = ws[f"{col}{r}"]
        c._style = copy.copy(tpl_cells[col])
    ws[f"A{r}"].fill = PatternFill(fill_type="solid",
                                   start_color=STAGE_FILL[stage],
                                   end_color=STAGE_FILL[stage])
ws.column_dimensions["A"].width = 22
ws.column_dimensions["B"].width = 16

wb.calculation.fullCalcOnLoad = True  # 打开时重算，保证公式结果始终最新
wb.save(DST)


# ---------------------------------------------------------- 写入公式缓存值
def _pct(num, den):
    """对齐 Excel TEXT(x,"0.0%") 的显示。"""
    return f"{num / max(1, den) * 100:.1f}%"


def _inject(xml, ref, typ, val):
    """给公式单元格补 <v> 缓存值，便于无公式引擎的预览器正确显示。"""
    m = re.search(r'<c r="%s"([^>]*)>(<f>.*?</f>)(?:<v>.*?</v>)?</c>' % re.escape(ref), xml, re.S)
    if not m:
        return xml, False
    attrs = re.sub(r'\st="[^"]*"', "", m.group(1))
    if typ == "str":
        attrs += ' t="str"'
    new = '<c r="%s"%s>%s<v>%s</v></c>' % (ref, attrs, m.group(2), escape(val))
    return xml[:m.start()] + new + xml[m.end():], True


def patch_cached(path, sheet_name, cached):
    with zipfile.ZipFile(path) as zin:
        data = {n: zin.read(n) for n in zin.namelist()}
    wbxml = data["xl/workbook.xml"].decode("utf-8")
    rels = data["xl/_rels/workbook.xml.rels"].decode("utf-8")
    rid = re.search(r'<sheet[^>]*name="%s"[^>]*r:id="([^"]+)"' % re.escape(sheet_name),
                    wbxml).group(1)
    rel = re.search(r'<Relationship[^>]*Id="%s"[^>]*/>' % re.escape(rid), rels).group(0)
    target = re.search(r'Target="([^"]+)"', rel).group(1).lstrip("/")
    if not target.startswith("xl/"):
        target = "xl/" + target
    xml = data[target].decode("utf-8")
    hit = 0
    for ref, (typ, val) in cached.items():
        xml, ok = _inject(xml, ref, typ, val)
        hit += ok
    data[target] = xml.encode("utf-8")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for n, blob in data.items():
            zout.writestr(n, blob)
    return hit


total = len(rows)
cached = {}
for i, s in enumerate(STATUSES):
    r = 4 + i
    n = sum(1 for x in rows if x[3] == s)
    cached[f"B{r}"] = ("n", str(n))
    cached[f"C{r}"] = ("str", _pct(n, total))
cached["B8"] = ("n", str(total))
for i, (stage, _rng, tasks) in enumerate(STAGES):
    r = 12 + i
    cnt = len(tasks)
    done = sum(1 for num, _ in tasks if num not in IN_PROGRESS)
    ing = cnt - done
    cached[f"C{r}"] = ("n", str(cnt))
    cached[f"D{r}"] = ("n", str(done))
    cached[f"E{r}"] = ("n", "0")
    cached[f"F{r}"] = ("n", str(ing))
    cached[f"G{r}"] = ("n", "0")
    cached[f"H{r}"] = ("str", _pct(done, cnt))

n = patch_cached(DST, "进度统计", cached)
print("saved:", DST)
print("cached values injected:", n, "/", len(cached))
print("tasks:", total,
      "| 已测试:", sum(1 for r in rows if r[3] == "已测试"),
      "| 开发中:", sum(1 for r in rows if r[3] == "开发中"))
