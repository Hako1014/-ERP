"""有界内容验证及私有本地存储。不执行文件内容，不生成公共链接。"""

import io
import re
import unicodedata
import warnings
from pathlib import Path, PurePosixPath
from uuid import uuid4
from zipfile import ZipFile

from defusedxml.ElementTree import fromstring
from PIL import Image, ImageOps
from pypdf import PdfReader

from app.core.config import get_settings
from app.core.errors import BizException, ErrorCode

MEDIA = {".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png", ".webp": "image/webp",
         ".pdf": "application/pdf", ".txt": "text/plain", ".csv": "text/csv",
         ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}
MIB = 1024 * 1024


def invalid(message):
    raise BizException(ErrorCode.VALIDATION_ERROR, message)


def policy():
    s = get_settings()
    return {"extensions": list(MEDIA), "max_file_bytes": s.assistant_file_max_mb * MIB,
            "max_files": s.assistant_files_per_message, "max_total_bytes": s.assistant_files_total_mb * MIB,
            "owner_quota_bytes": s.assistant_files_owner_quota_mb * MIB,
            "max_image_pixels": s.assistant_image_max_pixels, "max_pdf_pages": s.assistant_pdf_max_pages,
            "max_zip_entries": s.assistant_zip_max_entries, "max_uncompressed_bytes": s.assistant_zip_max_mb * MIB,
            "orphan_hours": s.assistant_orphan_hours, "deleted_retention_days": s.assistant_deleted_retention_days}


def safe_name(name):
    name = unicodedata.normalize("NFKC", name or "")
    name = name.replace("\\", "/").rsplit("/", 1)[-1]
    name = "".join(c for c in name if not unicodedata.category(c).startswith("C"))
    name = re.sub(r'[<>:"/\\|?*]', "_", name).strip(" .")
    ext = Path(name).suffix.lower()
    if ext not in MEDIA:
        invalid("仅支持 JPEG、PNG、WebP、PDF、TXT、CSV、XLSX，禁止宏及可执行文件")
    return name[:150 - len(ext)] + ext if len(name) > 150 else name


def decode_text(data):
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            value = data.decode(encoding)
            if any(ord(c) < 32 and c not in "\r\n\t" for c in value):
                invalid("文本包含二进制或不可用控制字符")
            return value
        except UnicodeDecodeError:
            pass
    invalid("文本编码不支持，请使用 UTF-8 或 GB18030")


def validate(data, name):
    name = safe_name(name)
    ext, limits = Path(name).suffix.lower(), policy()
    if not data or len(data) > limits["max_file_bytes"]:
        invalid(f"文件为空或超过 {get_settings().assistant_file_max_mb} MiB 限制")
    try:
        if ext in (".jpg", ".jpeg", ".png", ".webp"):
            with warnings.catch_warnings():
                warnings.simplefilter("error", Image.DecompressionBombWarning)
                with Image.open(io.BytesIO(data)) as image:
                    expected = {".jpg": "JPEG", ".jpeg": "JPEG", ".png": "PNG", ".webp": "WEBP"}[ext]
                    if image.format != expected or image.width * image.height > limits["max_image_pixels"]:
                        invalid("图片真实格式不符或像素超过限制")
                    if getattr(image, "n_frames", 1) != 1:
                        invalid("暂不支持多帧图片，请上传静态图片")
                    image.load()
                    clean = ImageOps.exif_transpose(image).convert("RGB" if expected == "JPEG" else "RGBA")
                    # 新图仅复制像素；不保留 EXIF、ICC、注释等上传者元数据。
                    output = io.BytesIO()
                    Image.frombytes(clean.mode, clean.size, clean.tobytes()).save(output, format=expected)
                    data = output.getvalue()
        elif ext == ".pdf":
            if not data.startswith(b"%PDF-"):
                invalid("PDF 真实格式不符")
            reader = PdfReader(io.BytesIO(data), strict=True)
            if reader.is_encrypted:
                invalid("不支持加密 PDF，请解密后上传")
            if not 0 < len(reader.pages) <= limits["max_pdf_pages"]:
                invalid(f"PDF 页数须为 1～{limits['max_pdf_pages']} 页，请拆分后上传")
            # 不执行/渲染 PDF 的脚本、附件或外部动作；预览仅返回元数据。
        elif ext in (".txt", ".csv"):
            decode_text(data)
        else:
            validate_xlsx(data, limits)
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
            try:
                if not workbook.worksheets:
                    invalid("XLSX 缺少可读取工作表")
                for sheet in workbook.worksheets:
                    next(sheet.iter_rows(min_row=1, max_row=1, max_col=10, values_only=True), None)
            finally:
                workbook.close()
    except BizException:
        raise
    except Exception:
        # 不回显解析器中的路径、文件正文或堆栈。
        invalid("文件损坏、加密或内容无法安全验证")
    if len(data) > limits["max_file_bytes"]:
        invalid("图片安全转码后超过单文件大小限制，请压缩后上传")
    return data, name, MEDIA[ext]


def validate_xlsx(data, limits):
    with ZipFile(io.BytesIO(data)) as archive:
        entries = archive.infolist()
        if len(entries) > limits["max_zip_entries"] or sum(i.file_size for i in entries) > limits["max_uncompressed_bytes"]:
            invalid("XLSX 解压大小或文件数量超限")
        names = [i.filename for i in entries]
        if len(set(names)) != len(names) or not {"[Content_Types].xml", "xl/workbook.xml"}.issubset(names):
            invalid("不是有效 XLSX 工作簿")
        for item in entries:
            path = PurePosixPath(item.filename)
            lower = item.filename.lower()
            if path.is_absolute() or ".." in path.parts or "\\" in lower or item.flag_bits & 1:
                invalid("XLSX 路径或加密内容不受支持")
            if item.file_size > max(item.compress_size, 1) * 200:
                invalid("XLSX 压缩比例异常，请拆分文件")
            if any(word in lower for word in ("vbaproject", "macrosheets", "embeddings/", "externallinks/", "activex/")):
                invalid("禁止含宏、嵌入对象或外部链接的 XLSX")
            if lower.endswith((".xml", ".rels")):
                document = fromstring(archive.read(item))
                for node in document.iter():
                    if node.attrib.get("TargetMode", "").lower() == "external":
                        invalid("禁止含外部资源链接的 XLSX")
                    if "macroenabled" in node.attrib.get("ContentType", "").lower():
                        invalid("禁止宏工作簿")


class PrivateAttachmentStorage:
    def __init__(self):
        settings = get_settings()
        self.root = Path(settings.assistant_private_dir).resolve()
        public = Path(settings.upload_dir).resolve()
        if self.root == public or self.root.is_relative_to(public):
            raise BizException(ErrorCode.INTERNAL_ERROR, "私有附件目录配置不安全")

    def path(self, key):
        if not re.fullmatch(r"[0-9a-f]{64}", key or ""):
            raise BizException(ErrorCode.INTERNAL_ERROR, "私有附件存储键无效")
        target = self.root / key
        if target.is_symlink() or target.resolve().parent != self.root:
            raise BizException(ErrorCode.INTERNAL_ERROR, "私有附件路径不安全")
        return target

    def save(self, data, *, key=None):
        key = key or uuid4().hex + uuid4().hex
        target = self.path(key)
        created = False
        try:
            self.root.mkdir(parents=True, exist_ok=True)
            with target.open("xb") as output:
                created = True
                output.write(data)
        except OSError:
            if created:
                target.unlink(missing_ok=True)
            raise BizException(ErrorCode.INTERNAL_ERROR, "附件保存失败，可能空间不足，请稍后重试") from None
        return key

    def read(self, key):
        try:
            target = self.path(key)
            if target.stat().st_size > policy()["max_file_bytes"]:
                invalid("附件已超过当前大小限制")
            return target.read_bytes()
        except OSError:
            raise BizException(ErrorCode.NOT_FOUND, "附件文件暂不可用") from None

    def remove(self, key):
        self.path(key).unlink(missing_ok=True)


def preview(data, media_type):
    """仅安全文本/表格与图片；PDF 不启用浏览器插件或脚本。预览截取明确标记。"""
    if media_type == "application/pdf":
        return {"kind": "metadata", "pages": len(PdfReader(io.BytesIO(data)).pages),
                "note": "PDF 仅预览页数；可下载核对。图文识别在后续任务实现。", "truncated": False}
    if media_type in ("text/plain", "text/csv"):
        value = decode_text(data)
        return {"kind": "text", "text": value[:12000], "truncated": len(value) > 12000,
                "note": "按纯文本显示，不执行 HTML 或 CSV 公式"}
    if media_type == MEDIA[".xlsx"]:
        from openpyxl import load_workbook
        validate_xlsx(data, policy())
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        try:
            sheet = workbook.worksheets[0]
            rows = [[str(cell)[:300] if cell is not None else "" for cell in row]
                    for row in sheet.iter_rows(min_row=1, max_row=20, max_col=10, values_only=True)]
            return {"kind": "table", "rows": rows, "sheet": sheet.title, "truncated": True,
                    "note": "仅预览首表前 20 行、10 列，单元格最多 300 字；公式仅作为文本，不计算"}
        finally:
            workbook.close()
    return {"kind": "image", "truncated": False}
