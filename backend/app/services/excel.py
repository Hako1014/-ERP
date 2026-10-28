from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter


@dataclass
class Column:
    field: str
    header: str
    required: bool = False
    type: str = "str"  # str / int / decimal / date
    width: int = 16


@dataclass
class SheetSpec:
    name: str
    columns: list[Column]
    sheet_title: str = "Sheet1"


@dataclass
class ParseResult:
    rows: list[dict[str, Any]] = field(default_factory=list)
    errors: list[dict[str, Any]] = field(default_factory=list)


TEMPLATE_REGISTRY: dict[str, SheetSpec] = {}


def register_template(name: str, spec: SheetSpec) -> None:
    TEMPLATE_REGISTRY[name] = spec


def _coerce(value: Any, type_name: str) -> Any:
    if value is None:
        return None
    if isinstance(value, str) and value.strip() == "":
        return None
    if type_name == "str":
        return str(value).strip()
    if type_name == "int":
        if isinstance(value, bool):
            raise ValueError("不是整数")
        if isinstance(value, float) and not value.is_integer():
            raise ValueError("不是整数")
        return int(value)
    if type_name == "decimal":
        return Decimal(str(value))
    if type_name == "date":
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        return datetime.strptime(str(value), "%Y-%m-%d").date()
    return value


def _new_workbook(spec: SheetSpec, rows: list[dict[str, Any]] | None = None) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = spec.sheet_title
    ws.append([c.header for c in spec.columns])
    for i, c in enumerate(spec.columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = c.width
    if rows:
        for row in rows:
            ws.append([row.get(c.field) for c in spec.columns])
    return wb


def _to_bytes(wb: Workbook) -> bytes:
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_template(spec: SheetSpec) -> bytes:
    return _to_bytes(_new_workbook(spec))


def export_data(spec: SheetSpec, rows: list[dict[str, Any]]) -> bytes:
    return _to_bytes(_new_workbook(spec, rows))


def parse_data(file_bytes: bytes, spec: SheetSpec, max_rows: int = 10000) -> ParseResult:
    try:
        wb = load_workbook(BytesIO(file_bytes), data_only=True)
    except Exception:
        return ParseResult([], [{"row": 0, "field": "", "message": "无法读取 Excel 文件"}])

    ws = wb[spec.sheet_title] if spec.sheet_title in wb.sheetnames else wb.active
    raw_rows = list(ws.iter_rows(values_only=True))
    if not raw_rows:
        return ParseResult([], [{"row": 0, "field": "", "message": "文件为空"}])
    if len(raw_rows) - 1 > max_rows:
        return ParseResult([], [{"row": 0, "field": "", "message": f"数据行数超过上限 {max_rows}"}])

    header = [str(c).strip() if c is not None else "" for c in raw_rows[0]]
    expected = [c.header for c in spec.columns]
    if header[: len(expected)] != expected:
        return ParseResult([], [{"row": 1, "field": "", "message": "表头不正确"}])

    result = ParseResult()
    for idx, row in enumerate(raw_rows[1:], start=2):
        if all(c is None or (isinstance(c, str) and c.strip() == "") for c in row):
            continue
        data: dict[str, Any] = {}
        row_errors: list[dict[str, Any]] = []
        for i, col in enumerate(spec.columns):
            raw = row[i] if i < len(row) else None
            is_empty = raw is None or (isinstance(raw, str) and raw.strip() == "")
            if is_empty:
                if col.required:
                    row_errors.append(
                        {"row": idx, "field": col.field, "message": f"{col.header} 为必填项"}
                    )
                else:
                    data[col.field] = None
                continue
            try:
                data[col.field] = _coerce(raw, col.type)
            except (ValueError, InvalidOperation, TypeError):
                row_errors.append(
                    {"row": idx, "field": col.field, "message": f"{col.header} 格式不正确"}
                )
        if row_errors:
            result.errors.extend(row_errors)
        else:
            result.rows.append(data)
    return result
