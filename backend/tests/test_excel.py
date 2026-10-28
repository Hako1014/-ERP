from datetime import datetime
from decimal import Decimal
from io import BytesIO

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.api.common import router
from app.services.excel import (
    Column,
    SheetSpec,
    build_template,
    export_data,
    parse_data,
    register_template,
)

SPEC = SheetSpec(
    name="sample",
    columns=[
        Column("name", "名称", required=True, type="str"),
        Column("qty", "数量", required=True, type="int"),
        Column("price", "单价", type="decimal"),
        Column("date", "日期", type="date"),
    ],
)


def _make_xlsx(rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_build_template():
    wb = load_workbook(BytesIO(build_template(SPEC)))
    assert [c.value for c in wb.active[1]] == ["名称", "数量", "单价", "日期"]


def test_export_data():
    wb = load_workbook(BytesIO(export_data(SPEC, [{"name": "苹果", "qty": 3, "price": Decimal("1.5")}])))
    ws = wb.active
    assert [c.value for c in ws[1]] == ["名称", "数量", "单价", "日期"]
    row2 = [c.value for c in ws[2]]
    assert row2[0] == "苹果"
    assert row2[1] == 3
    assert round(float(row2[2]), 2) == 1.5


def test_parse_valid():
    data = _make_xlsx([
        ["名称", "数量", "单价", "日期"],
        ["苹果", 3, "1.5", datetime(2026, 9, 23)],
    ])
    result = parse_data(data, SPEC)
    assert result.errors == []
    assert result.rows == [
        {"name": "苹果", "qty": 3, "price": Decimal("1.5"), "date": datetime(2026, 9, 23).date()}
    ]


def test_parse_invalid_rows():
    data = _make_xlsx([
        ["名称", "数量", "单价", "日期"],
        ["", 3, "1.5", None],
        ["香蕉", "abc", "1.5", None],
        ["梨", 2, "1.5", "bad-date"],
    ])
    result = parse_data(data, SPEC)
    assert result.rows == []
    assert [e["row"] for e in result.errors] == [2, 3, 4]
    assert result.errors[0]["field"] == "name"
    assert result.errors[1]["field"] == "qty"
    assert result.errors[2]["field"] == "date"


def test_parse_wrong_header():
    data = _make_xlsx([["错误表头", "数量", "单价", "日期"], ["苹果", 3, "1.5", None]])
    result = parse_data(data, SPEC)
    assert result.rows == []
    assert result.errors[0]["message"] == "表头不正确"


def test_template_endpoint():
    from app.api.deps import get_current_staff
    from app.core.handlers import register_exception_handlers

    register_template("sample", SPEC)
    app = FastAPI()
    register_exception_handlers(app)
    app.include_router(router)
    client = TestClient(app)
    assert client.get("/common/excel-template/sample").status_code == 401
    app.dependency_overrides[get_current_staff] = lambda: object()

    resp = client.get("/common/excel-template/sample")
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert resp.content[:2] == b"PK"
