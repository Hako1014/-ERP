from decimal import Decimal
from io import BytesIO

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

import app.models  # noqa: F401
from app.api.stock_check import router as stock_check_router
from app.core.data_scope import initialize_scope_policy
from app.core.errors import BizException
from app.core.handlers import register_exception_handlers
from app.core.security import create_access_token
from app.db.base import Base
from app.db.session import get_db
from app.models.basic import Unit
from app.models.finance import Voucher
from app.models.org import Company, Staff
from app.models.product import Product, ProductSku
from app.models.rbac import Role, StaffRole
from app.models.stock import (
    StockCheck,
    StockCheckItem,
    StockFlow,
    StockIn,
    StockInItem,
    StockOut,
    StockOutItem,
)
from app.models.warehouse import Location, StockLocation, Warehouse
from app.services import (
    approval_service,
    stock_check_service,
    stock_in_service,
    stock_query_service,
    voucher_service,
)

XLSX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)


def _session_factory(pool: bool = False):
    kwargs = (
        {"connect_args": {"check_same_thread": False}, "poolclass": StaticPool}
        if pool
        else {}
    )
    engine = create_engine("sqlite:///:memory:", future=True, **kwargs)
    Base.metadata.create_all(engine)
    return sessionmaker(bind=engine, autoflush=False, expire_on_commit=False)


def _seed(db) -> None:
    db.add(Company(id=1, full_name="示例公司"))
    db.add_all(
        [
            Staff(id=1, company_id=1, name="仓库", phone="13900000001"),
            Staff(id=2, company_id=1, name="审核员", phone="13900000002"),
            Staff(id=3, company_id=1, name="质检", phone="13900000003"),
        ]
    )
    db.add(Unit(id=1, code="PCS", name="件"))
    db.add_all(
        [
            Product(
                id=1,
                name="主商品",
                code="P0001",
                main_unit_id=1,
                cost_price=Decimal("10.00"),
            ),
            Product(
                id=2,
                name="配件商品",
                code="P0002",
                main_unit_id=1,
                cost_price=Decimal("5.00"),
            ),
        ]
    )
    db.add_all(
        [
            ProductSku(
                id=1,
                product_id=1,
                spec_name="默认",
                spec_code="P0001-01",
                barcode="690000000001",
                cost_price=Decimal("10.00"),
                is_active=True,
            ),
            ProductSku(
                id=2,
                product_id=2,
                spec_name="默认",
                spec_code="P0002-01",
                barcode="690000000002",
                cost_price=Decimal("5.00"),
                is_active=True,
            ),
        ]
    )
    db.add_all(
        [
            Warehouse(
                id=1,
                code="WH01",
                name="主仓",
                warehouse_type="FINISHED",
                is_enabled=True,
                enable_location=True,
                lock_type="NONE",
            ),
            Warehouse(
                id=2,
                code="WH02",
                name="次品仓",
                warehouse_type="DEFECTIVE",
                is_enabled=True,
                enable_location=True,
                lock_type="NONE",
            ),
            Warehouse(
                id=3,
                code="VWH",
                name="虚拟仓",
                warehouse_type="FINISHED",
                is_enabled=True,
                enable_location=True,
                is_virtual=True,
                lock_type="NONE",
            ),
        ]
    )
    db.add_all(
        [
            Location(id=1, warehouse_id=1, code="WH01-A01", is_enabled=True),
            Location(id=2, warehouse_id=1, code="WH01-A02", is_enabled=True),
            Location(id=3, warehouse_id=2, code="WH02-A01", is_enabled=True),
            Location(id=4, warehouse_id=3, code="VWH-A01", is_enabled=True),
        ]
    )
    db.add_all(
        [
            StockLocation(
                id=1,
                location_id=1,
                product_sku_id=1,
                qty=Decimal("10.000"),
                status="NORMAL",
            ),
            StockLocation(
                id=2,
                location_id=2,
                product_sku_id=1,
                qty=Decimal("5.000"),
                status="NORMAL",
            ),
            StockLocation(
                id=3,
                location_id=1,
                product_sku_id=2,
                qty=Decimal("3.000"),
                status="NORMAL",
            ),
            StockLocation(
                id=4,
                location_id=3,
                product_sku_id=1,
                qty=Decimal("2.000"),
                status="DEFECTIVE",
            ),
        ]
    )
    db.add_all(
        [
            Role(id=1, company_id=1, code="WAREHOUSE", name="仓库"),
            Role(id=2, company_id=1, code="ADMIN", name="管理员"),
            Role(id=3, company_id=1, code="QUALITY", name="质检"),
        ]
    )
    db.flush()
    db.add_all(
        [
            StaffRole(staff_id=1, role_id=1),
            StaffRole(staff_id=2, role_id=2),
            StaffRole(staff_id=3, role_id=3),
        ]
    )
    db.commit()
    approval_service.create_flow(
        db,
        name="库存盘点审批",
        doc_type=stock_check_service.DOC_TYPE,
        approvers=[2],
        actor=db.get(Staff, 1),
    )
    initialize_scope_policy(db, db.get(Staff, 1))


def _patch_serial(monkeypatch) -> None:
    counters: dict[str, int] = {}

    def fake_serial(prefix, dt=None):
        code = prefix.value if hasattr(prefix, "value") else str(prefix)
        counters[code] = counters.get(code, 0) + 1
        return f"{code}20260923-{counters[code]:04d}"

    monkeypatch.setattr(
        stock_check_service,
        "generate_serial",
        fake_serial,
    )
    monkeypatch.setattr(
        stock_in_service,
        "generate_serial",
        fake_serial,
    )


def _create_check(
    db,
    *,
    items: list[dict] | None = None,
    lock_warehouse: bool = True,
):
    return stock_check_service.create_stock_check(
        db,
        data={
            "warehouse_id": 1,
            "lock_warehouse": lock_warehouse,
            "remark": "测试盘点",
            "items": items,
        },
        staff=db.get(Staff, 1),
    )


def _full_actual_items() -> list[dict]:
    return [
        {
            "product_id": 1,
            "sku_id": 1,
            "location_id": 1,
            "actual_qty": "12.000",
        },
        {
            "product_id": 1,
            "sku_id": 1,
            "location_id": 2,
            "actual_qty": "4.000",
        },
        {
            "product_id": 2,
            "sku_id": 2,
            "location_id": 1,
            "actual_qty": "3.000",
        },
    ]


def test_stock_check_approval_updates_locations_and_generates_documents(
    monkeypatch,
):
    _patch_serial(monkeypatch)
    Session = _session_factory()
    with Session() as db:
        _seed(db)
        record = _create_check(db)
        assert record.check_no.startswith("PDD20260923-")
        assert record.lock_warehouse is True
        assert db.get(Warehouse, 1).lock_type == "CHECK"

        detail = stock_check_service.stock_check_dict(
            db,
            record,
            with_items=True,
        )
        assert detail["item_count"] == 3
        assert detail["warehouse_locked"] is True
        assert all(item["actual_qty"] is None for item in detail["items"])

        with pytest.raises(BizException) as incomplete:
            stock_check_service.submit_stock_check(
                db,
                check_id=record.id,
                staff=db.get(Staff, 1),
            )
        assert "未录入盘点实存" in incomplete.value.message

        stock_check_service.update_stock_check(
            db,
            check_id=record.id,
            data={"items": _full_actual_items()},
            staff=db.get(Staff, 1),
        )
        stock_check_service.submit_stock_check(
            db,
            check_id=record.id,
            staff=db.get(Staff, 1),
        )
        approved = stock_check_service.approve_stock_check(
            db,
            check_id=record.id,
            staff=db.get(Staff, 2),
            comment="同意",
        )

        assert approved.status == "APPROVED"
        assert approved.audit_status == "APPROVED"
        assert db.get(Warehouse, 1).lock_type == "NONE"
        assert db.get(StockLocation, 1).qty == Decimal("12.000")
        assert db.get(StockLocation, 2).qty == Decimal("4.000")
        assert db.get(StockLocation, 3).qty == Decimal("3.000")

        flows = list(
            db.scalars(
                select(StockFlow)
                .where(StockFlow.source_no == record.check_no)
                .order_by(StockFlow.id)
            )
        )
        assert [(row.biz_type, row.qty) for row in flows] == [
            ("CHECK_IN", Decimal("2.000")),
            ("CHECK_OUT", Decimal("1.000")),
        ]
        assert [row.total_cost for row in flows] == [
            Decimal("20.00"),
            Decimal("10.00"),
        ]

        stock_in = db.scalar(
            select(StockIn).where(StockIn.remark == f"盘点盘盈:{record.check_no}")
        )
        stock_out = db.scalar(
            select(StockOut).where(StockOut.remark == f"盘点盘亏:{record.check_no}")
        )
        assert stock_in is not None
        assert stock_in.stock_in_type == "CHECK"
        assert stock_in.status == "COMPLETED"
        assert stock_in.total_qty == Decimal("2.000")
        assert (
            db.scalar(
                select(func.count(StockInItem.id)).where(
                    StockInItem.stock_in_id == stock_in.id
                )
            )
            == 1
        )
        voucher = db.scalar(
            select(Voucher).where(
                Voucher.reference == f"采购入库单:{stock_in.stock_in_no}"
            )
        )
        assert voucher is not None
        voucher_detail = voucher_service.voucher_dict(db, voucher)
        assert [item["subject_code"] for item in voucher_detail["items"]] == [
            "1405",
            "1901",
        ]
        assert voucher.total_debit == Decimal("20.00")
        assert voucher.total_credit == Decimal("20.00")
        assert stock_out is not None
        assert stock_out.stock_out_type == "CHECK"
        assert stock_out.off_shelf_status == "DONE"
        assert stock_out.total_qty == Decimal("1.000")
        assert (
            db.scalar(
                select(func.count(StockOutItem.id)).where(
                    StockOutItem.stock_out_id == stock_out.id
                )
            )
            == 1
        )


def test_stock_check_lock_release_on_update_delete_reject_and_terminate(
    monkeypatch,
):
    _patch_serial(monkeypatch)
    Session = _session_factory()
    with Session() as db:
        _seed(db)
        staff = db.get(Staff, 1)
        record = _create_check(
            db,
            items=[
                {
                    "product_id": 1,
                    "sku_id": 1,
                    "location_id": 1,
                    "actual_qty": "10.000",
                }
            ],
        )
        with pytest.raises(BizException) as duplicate:
            _create_check(db, items=[])
        assert "进行中的盘点单" in duplicate.value.message

        updated = stock_check_service.update_stock_check(
            db,
            check_id=record.id,
            data={"lock_warehouse": False},
            staff=staff,
        )
        assert updated.lock_warehouse is False
        assert db.get(Warehouse, 1).lock_type == "NONE"

        stock_check_service.delete_stock_check(
            db,
            check_id=record.id,
            staff=staff,
        )
        assert db.get(StockCheck, record.id) is None
        assert (
            db.scalar(
                select(func.count(StockCheckItem.id)).where(
                    StockCheckItem.check_id == record.id
                )
            )
            == 0
        )

        rejected = _create_check(
            db,
            items=[
                {
                    "product_id": 1,
                    "sku_id": 1,
                    "location_id": 1,
                    "actual_qty": "9.000",
                }
            ],
        )
        stock_check_service.submit_stock_check(
            db,
            check_id=rejected.id,
            staff=staff,
        )
        rejected = stock_check_service.reject_stock_check(
            db,
            check_id=rejected.id,
            staff=db.get(Staff, 2),
            comment="重盘",
        )
        assert rejected.status == "REJECTED"
        assert db.get(Warehouse, 1).lock_type == "NONE"

        terminated = stock_check_service.terminate_stock_check(
            db,
            check_id=rejected.id,
            staff=staff,
            reason="不再盘点",
        )
        assert terminated.status == "TERMINATED"


def test_stock_check_validates_location_snapshot_and_warehouse_type(
    monkeypatch,
):
    _patch_serial(monkeypatch)
    Session = _session_factory()
    with Session() as db:
        _seed(db)
        staff = db.get(Staff, 1)

        with pytest.raises(BizException) as foreign:
            _create_check(
                db,
                items=[
                    {
                        "product_id": 1,
                        "sku_id": 1,
                        "location_id": 3,
                        "actual_qty": "1.000",
                    }
                ],
            )
        assert "不属于当前仓库" in foreign.value.message

        with pytest.raises(BizException) as virtual:
            stock_check_service.create_stock_check(
                db,
                data={
                    "warehouse_id": 3,
                    "lock_warehouse": False,
                    "items": [],
                },
                staff=staff,
            )
        assert "虚拟仓" in virtual.value.message

        record = _create_check(
            db,
            lock_warehouse=False,
            items=[
                {
                    "product_id": 1,
                    "sku_id": 1,
                    "location_id": 1,
                    "actual_qty": "8.000",
                }
            ],
        )
        db.get(StockLocation, 1).qty = Decimal("9.000")
        db.commit()
        with pytest.raises(BizException) as changed:
            stock_check_service.submit_stock_check(
                db,
                check_id=record.id,
                staff=staff,
            )
        assert "库存发生变化" in changed.value.message
        assert db.get(StockLocation, 1).qty == Decimal("9.000")


def _make_import_workbook(rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "库存盘点"
    ws.append(
        [
            "商品编码",
            "商品名称",
            "规格编码",
            "规格名称",
            "所在库位",
            "条形码",
            "现有库存",
            "盘点实存",
        ]
    )
    for row in rows:
        ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_stock_check_excel_export_and_import_include_location(monkeypatch):
    _patch_serial(monkeypatch)
    Session = _session_factory()
    with Session() as db:
        _seed(db)
        record = _create_check(
            db,
            items=[
                {
                    "product_id": 1,
                    "sku_id": 1,
                    "location_id": 1,
                    "actual_qty": None,
                }
            ],
        )

        exported = stock_check_service.export_stock_check(
            db,
            check_id=record.id,
        )
        ws = load_workbook(BytesIO(exported)).active
        assert [cell.value for cell in ws[1]][4] == "所在库位"
        assert ws.cell(row=2, column=5).value == "WH01-A01"
        assert ws.cell(row=2, column=7).value == 10

        imported = stock_check_service.import_stock_check_result(
            db,
            check_id=record.id,
            file_bytes=_make_import_workbook(
                [["P0001", "主商品", "P0001-01", "默认", "WH01-A01", "", 10, 11]]
            ),
            staff=db.get(Staff, 1),
        )
        assert imported["updated"] == 1
        item = db.scalar(
            select(StockCheckItem).where(
                StockCheckItem.check_id == record.id
            )
        )
        assert item.actual_qty == Decimal("11.000")
        assert stock_check_service.stock_check_dict(
            db,
            record,
        )["has_imported_result"] is True

        invalid = stock_check_service.import_stock_check_result(
            db,
            check_id=record.id,
            file_bytes=_make_import_workbook(
                [["P0001", "主商品", "P0001-01", "默认", "WH99", "", 10, 8]]
            ),
            staff=db.get(Staff, 1),
        )
        assert invalid["updated"] == 0
        assert "库位不存在" in invalid["errors"][0]["message"]
        assert db.get(StockCheckItem, item.id).actual_qty == Decimal("11.000")


def test_stock_query_services_respect_stock_check_lock(monkeypatch):
    _patch_serial(monkeypatch)
    Session = _session_factory()
    with Session() as db:
        _seed(db)
        _create_check(
            db,
            items=[
                {
                    "product_id": 1,
                    "sku_id": 1,
                    "location_id": 1,
                    "actual_qty": "10.000",
                }
            ],
        )
        with pytest.raises(BizException) as locked:
            stock_query_service.manual_out(
                db,
                data={
                    "warehouse_id": 1,
                    "location_id": 1,
                    "product_id": 1,
                    "sku_id": 1,
                    "qty": "1.000",
                },
                staff=db.get(Staff, 1),
            )
        assert "锁定" in locked.value.message


def _make_api():
    Session = _session_factory(pool=True)

    def override_get_db():
        db = Session()
        try:
            yield db
        finally:
            db.close()

    app = FastAPI()
    register_exception_handlers(app)
    app.include_router(stock_check_router)
    app.dependency_overrides[get_db] = override_get_db
    return app, Session


def _auth(staff_id: int, name: str) -> dict:
    return {"Authorization": f"Bearer {create_access_token(staff_id, name, [])}"}


def test_api_stock_check_permissions_and_full_approval(monkeypatch):
    _patch_serial(monkeypatch)
    app, Session = _make_api()
    with Session() as db:
        _seed(db)

    client = TestClient(app)
    warehouse_headers = _auth(1, "仓库")
    admin_headers = _auth(2, "审核员")
    quality_headers = _auth(3, "质检")

    readable = client.get("/stock-checks", headers=quality_headers)
    assert readable.status_code == 200

    denied = client.post(
        "/stock-checks",
        json={"warehouse_id": 1},
        headers=quality_headers,
    )
    assert denied.status_code == 403

    created = client.post(
        "/stock-checks",
        json={
            "warehouse_id": 1,
            "lock_warehouse": True,
            "items": [
                {
                    "product_id": 1,
                    "sku_id": 1,
                    "location_id": 1,
                    "actual_qty": "8.000",
                }
            ],
        },
        headers=warehouse_headers,
    )
    assert created.status_code == 200
    check_id = created.json()["data"]["id"]

    template = client.get(
        "/stock-checks/template",
        headers=quality_headers,
    )
    assert template.status_code == 200
    assert template.headers["content-type"].startswith(XLSX_MEDIA_TYPE)

    exported = client.get(
        f"/stock-checks/{check_id}/export",
        headers=quality_headers,
    )
    assert exported.status_code == 404
    exported = client.get(
        f"/stock-checks/{check_id}/export",
        headers=warehouse_headers,
    )
    assert exported.status_code == 200
    assert exported.headers["content-type"].startswith(XLSX_MEDIA_TYPE)

    submitted = client.post(
        f"/stock-checks/{check_id}/submit",
        headers=warehouse_headers,
    )
    assert submitted.status_code == 200
    assert submitted.json()["data"]["audit_status"] == "PENDING"

    denied_approve = client.post(
        f"/stock-checks/{check_id}/approve",
        json={"comment": "越权"},
        headers=quality_headers,
    )
    assert denied_approve.status_code == 403

    approved = client.post(
        f"/stock-checks/{check_id}/approve",
        json={"comment": "同意"},
        headers=admin_headers,
    )
    assert approved.status_code == 200
    assert approved.json()["data"]["status"] == "APPROVED"

    with Session() as db:
        assert db.get(StockLocation, 1).qty == Decimal("8.000")
        assert db.get(Warehouse, 1).lock_type == "NONE"
