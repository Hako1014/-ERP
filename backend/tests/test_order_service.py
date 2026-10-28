from datetime import datetime
from decimal import Decimal
from io import BytesIO

import pytest
from fastapi import FastAPI
from fastapi.encoders import jsonable_encoder
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

import app.models  # noqa: F401
from app.api.order import router as order_router
from app.core.errors import BizException
from app.core.handlers import register_exception_handlers
from app.core.security import create_access_token
from app.db.base import Base
from app.db.session import get_db
from app.models.basic import Unit
from app.models.org import Company, Staff
from app.models.platform import (
    AreaRule,
    FreightTemplate,
    Logistics,
    Platform,
    Shop,
)
from app.models.product import Product, ProductSku
from app.models.rbac import Role, StaffRole
from app.models.warehouse import Location, StockLocation, Warehouse
from app.services import approval_service, order_service
from app.services.excel import build_template


def _session_factory(pool: bool = False):
    kwargs = (
        {"connect_args": {"check_same_thread": False}, "poolclass": StaticPool}
        if pool
        else {}
    )
    engine = create_engine("sqlite:///:memory:", future=True, **kwargs)
    Base.metadata.create_all(engine)
    return sessionmaker(bind=engine, autoflush=False, expire_on_commit=False)


def _seed(db) -> None:
    db.add(Company(id=1, full_name="示例公司"))
    db.add_all(
        [
            Staff(id=1, company_id=1, name="仓库", phone="13900000001"),
            Staff(id=2, company_id=1, name="财务", phone="13900000002"),
            Staff(id=3, company_id=1, name="售后", phone="13900000003"),
        ]
    )
    db.add(Unit(id=1, code="PCS", name="件"))
    db.add_all(
        [
            Product(
                id=1,
                name="测试商品",
                code="P0001",
                main_unit_id=1,
                online_price=Decimal("10.00"),
                cost_price=Decimal("6.00"),
                weight=Decimal("0.500"),
            ),
            Product(
                id=2,
                name="另一商品",
                code="P0002",
                main_unit_id=1,
                online_price=Decimal("8.00"),
                cost_price=Decimal("5.00"),
                weight=Decimal("0.200"),
            ),
        ]
    )
    db.add_all(
        [
            ProductSku(
                id=1,
                product_id=1,
                spec_name="默认",
                spec_code="P0001-01",
                cost_price=Decimal("6.00"),
                is_active=True,
            ),
            ProductSku(
                id=2,
                product_id=2,
                spec_name="默认",
                spec_code="P0002-01",
                cost_price=Decimal("5.00"),
                is_active=True,
            ),
        ]
    )
    db.add_all(
        [
            Warehouse(
                id=1,
                code="WH01",
                name="主仓",
                warehouse_type="FINISHED",
                is_enabled=True,
                is_default=True,
                is_virtual=False,
                lock_type="NONE",
            ),
            Warehouse(
                id=2,
                code="WH02",
                name="区域仓",
                warehouse_type="FINISHED",
                is_enabled=True,
                is_default=False,
                is_virtual=False,
                lock_type="NONE",
                ship_regions="北京市",
            ),
        ]
    )
    db.add_all(
        [
            Location(id=1, warehouse_id=1, code="WH01-A01", is_enabled=True),
            Location(id=2, warehouse_id=2, code="WH02-A01", is_enabled=True),
        ]
    )
    db.add_all(
        [
            StockLocation(
                id=1,
                location_id=1,
                product_sku_id=1,
                qty=Decimal("10.000"),
                status="NORMAL",
            ),
            StockLocation(
                id=2,
                location_id=1,
                product_sku_id=2,
                qty=Decimal("10.000"),
                status="NORMAL",
            ),
        ]
    )
    db.add_all(
        [
            Platform(id=1, name="淘宝", code="TAOBAO"),
            Shop(
                id=1,
                platform_id=1,
                name="旗舰店",
                auto_sync_order=True,
            ),
        ]
    )
    db.add_all(
        [
            Logistics(
                id=1,
                name="经济快递",
                is_enabled=True,
                is_default=True,
            ),
            Logistics(
                id=2,
                name="标准快递",
                is_enabled=True,
                is_default=False,
            ),
        ]
    )
    db.add_all(
        [
            FreightTemplate(
                id=1,
                name="经济快递模板",
                logistics_id=1,
                charge_type="BY_COUNT",
                first_count=1,
                continue_count=1,
                base_fee=Decimal("5.00"),
                continue_fee=Decimal("1.00"),
                operation_fee=Decimal("0.00"),
                consumable_fee=Decimal("0.00"),
                ewaybill_fee=Decimal("0.00"),
                tax_fee=Decimal("0.00"),
                tax_point=Decimal("0.00"),
            ),
            FreightTemplate(
                id=2,
                name="标准快递模板",
                logistics_id=2,
                charge_type="BY_COUNT",
                first_count=1,
                continue_count=1,
                base_fee=Decimal("8.00"),
                continue_fee=Decimal("2.00"),
                operation_fee=Decimal("0.00"),
                consumable_fee=Decimal("0.00"),
                ewaybill_fee=Decimal("0.00"),
                tax_fee=Decimal("0.00"),
                tax_point=Decimal("0.00"),
            ),
        ]
    )
    db.flush()
    db.add_all(
        [
            AreaRule(
                id=1,
                logistics_id=1,
                province="浙江省",
                freight_template_id=1,
            ),
            AreaRule(
                id=2,
                logistics_id=2,
                province="浙江省",
                freight_template_id=2,
            ),
        ]
    )
    db.add_all(
        [
            Role(id=1, company_id=1, code="WAREHOUSE", name="仓库"),
            Role(id=2, company_id=1, code="FINANCE", name="财务"),
            Role(id=3, company_id=1, code="AFTER_SALE", name="售后"),
            Role(id=4, company_id=1, code="ADMIN", name="管理员"),
        ]
    )
    db.flush()
    db.add_all(
        [
            StaffRole(staff_id=1, role_id=1),
            StaffRole(staff_id=2, role_id=2),
            StaffRole(staff_id=3, role_id=3),
        ]
    )
    db.commit()


def _flow(db) -> None:
    approval_service.create_flow(
        db,
        name="订单审核流程",
        doc_type=order_service.DOC_TYPE,
        approvers=[2],
        actor=db.get(Staff, 2),
    )


def _order_payload(*, order_no: str = "TB20260923-001", **overrides) -> dict:
    payload = {
        "order_no": order_no,
        "shop_id": 1,
        "receiver_name": "张三",
        "receiver_phone": "13800000000",
        "receiver_province": "浙江省",
        "receiver_city": "杭州市",
        "receiver_address": "杭州市西湖区 1 号",
        "order_time": datetime(2026, 9, 23, 10, 0, 0),
        "pay_time": datetime(2026, 9, 23, 10, 1, 0),
        "pay_status": "PAID",
        "freight": "5.00",
        "platform_discount": "2.00",
        "seller_discount": "3.00",
        "items": [
            {
                "product_id": 1,
                "sku_id": 1,
                "qty": "2.000",
                "price": "10.00",
            }
        ],
    }
    payload.update(overrides)
    return payload


def test_create_order_recalculates_amount_weight_and_flags(monkeypatch):
    monkeypatch.setattr(
        order_service,
        "generate_serial",
        lambda prefix, dt=None: "DD20260923-0001",
    )
    Session = _session_factory()
    with Session() as db:
        _seed(db)
        order = order_service.create_order(
            db,
            data=_order_payload(
                order_no=None,
                flags=["DISTRIBUTION", "MERGED"],
            ),
            staff=db.get(Staff, 1),
        )
        detail = order_service.order_dict(db, order)
        assert order.order_no == "DD20260923-0001"
        assert detail["order_amount"] == Decimal("25.00")
        assert detail["paid_amount"] == Decimal("20.00")
        assert detail["total_qty"] == Decimal("2.000")
        assert detail["weight"] == Decimal("1.000")
        assert detail["flags"] == ["DISTRIBUTION", "MERGED"]


def test_approve_order_auto_assigns_warehouse_and_cheapest_logistics():
    Session = _session_factory()
    with Session() as db:
        _seed(db)
        _flow(db)
        order = order_service.create_order(
            db,
            data=_order_payload(),
            staff=db.get(Staff, 1),
        )

        order, checks = order_service.approve_order(
            db,
            order_id=order.id,
            staff=db.get(Staff, 2),
            comment="同意",
        )

        assert order.status == order_service.STATUS_PENDING
        assert order.audit_status == order_service.AUDIT_APPROVED
        assert order.warehouse_id == 1
        assert order.logistics_id == 1
        assert checks["auto_warehouse"] is True
        assert checks["auto_logistics"] is True
        assert checks["auto_submitted"] is True
        detail = order_service.order_dict(db, order, with_approval=True)
        assert detail["approval"]["state"]["status"] == "APPROVED"


def test_approve_order_rejects_stock_shortage_and_region_mismatch():
    Session = _session_factory()
    with Session() as db:
        _seed(db)
        _flow(db)
        finance = db.get(Staff, 2)

        shortage_order = order_service.create_order(
            db,
            data=_order_payload(
                order_no="TB20260923-002",
                warehouse_id=1,
                items=[
                    {
                        "product_id": 1,
                        "sku_id": 1,
                        "qty": "99.000",
                        "price": "10.00",
                    }
                ],
            ),
            staff=db.get(Staff, 1),
        )
        with pytest.raises(BizException) as shortage_exc:
            order_service.approve_order(
                db,
                order_id=shortage_order.id,
                staff=finance,
            )
        assert "库存不足" in shortage_exc.value.message

        region_order = order_service.create_order(
            db,
            data=_order_payload(
                order_no="TB20260923-003",
                warehouse_id=2,
                items=[
                    {
                        "product_id": 2,
                        "sku_id": 2,
                        "qty": "1.000",
                        "price": "8.00",
                    }
                ],
            ),
            staff=db.get(Staff, 1),
        )
        with pytest.raises(BizException) as region_exc:
            order_service.approve_order(
                db,
                order_id=region_order.id,
                staff=finance,
            )
        assert "收货区域" in region_exc.value.message


def test_set_warehouse_logistics_flags_and_logs():
    Session = _session_factory()
    with Session() as db:
        _seed(db)
        staff = db.get(Staff, 1)
        order = order_service.create_order(
            db, data=_order_payload(), staff=staff
        )

        order_service.set_order_warehouse(
            db, order_id=order.id, warehouse_id=2, staff=staff
        )
        order = order_service.set_order_logistics(
            db,
            order_id=order.id,
            logistics_id=2,
            tracking_no="SF123456",
            staff=staff,
        )
        order_service.set_order_flags(
            db,
            order_id=order.id,
            flags=["REFUND"],
            mode="ADD",
            staff=staff,
        )
        order = order_service.set_order_flags(
            db,
            order_id=order.id,
            flags=["REFUND"],
            mode="REMOVE",
            staff=staff,
        )

        assert order.warehouse_id == 2
        assert order.logistics_id == 2
        assert order.logistics_company == "标准快递"
        assert order.tracking_no == "SF123456"
        assert order.flags is None
        actions = {
            row["action"]
            for row in order_service.list_order_logs(
                db, order_id=order.id
            )
        }
        assert {"SET_WAREHOUSE", "SET_LOGISTICS", "SET_FLAGS"} <= actions


def test_submit_reject_approve_unaudit_and_operation_logs():
    Session = _session_factory()
    with Session() as db:
        _seed(db)
        _flow(db)
        warehouse = db.get(Staff, 1)
        finance = db.get(Staff, 2)
        order = order_service.create_order(
            db, data=_order_payload(), staff=warehouse
        )

        order = order_service.submit_order(
            db, order_id=order.id, staff=warehouse
        )
        assert order.audit_status == order_service.AUDIT_PENDING

        order = order_service.reject_order(
            db,
            order_id=order.id,
            staff=finance,
            comment="地址异常",
        )
        assert order.audit_status == order_service.AUDIT_REJECTED
        assert order.status == order_service.STATUS_NEW

        order = order_service.submit_order(
            db, order_id=order.id, staff=warehouse
        )
        order, _ = order_service.approve_order(
            db, order_id=order.id, staff=finance, comment="同意"
        )
        assert order.status == order_service.STATUS_PENDING

        order = order_service.unaudit_order(
            db,
            order_id=order.id,
            staff=finance,
            reason="重新核对运费",
        )
        assert order.status == order_service.STATUS_NEW
        assert order.audit_status == order_service.AUDIT_PENDING

        logs = order_service.list_order_logs(db, order_id=order.id)
        actions = {row["action"] for row in logs}
        assert {"SUBMIT", "REJECT", "APPROVE", "UNAUDIT"} <= actions
        assert any(row["source"] == "APPROVAL" for row in logs)


def test_fetch_orders_is_idempotent_and_skips_shipped_order():
    Session = _session_factory()
    with Session() as db:
        _seed(db)
        staff = db.get(Staff, 1)
        first = order_service.create_order(
            db, data=_order_payload(), staff=staff
        )
        shipped = order_service.create_order(
            db,
            data=_order_payload(order_no="TB20260923-002"),
            staff=staff,
        )
        shipped.status = order_service.STATUS_SHIPPED
        db.commit()

        payload = _order_payload()
        payload["items"][0]["qty"] = "3.000"
        result = order_service.fetch_orders(
            db,
            orders=[
                payload,
                _order_payload(order_no="TB20260923-002"),
            ],
            staff=staff,
        )

        assert result["created"] == 0
        assert result["updated"] == 1
        assert len(result["skipped"]) == 1
        assert result["skipped"][0]["order_no"] == shipped.order_no
        first_detail = order_service.order_dict(db, first)
        assert first_detail["total_qty"] == Decimal("3.000")
        assert db.get(type(first), first.id).status == order_service.STATUS_NEW


def test_platform_adapter_pull_and_celery_task_eager(monkeypatch):
    from app.integrations import platform_tasks
    from app.integrations.platform import MockPlatformAdapter
    from app.services import platform_sync_service

    Session = _session_factory(pool=True)
    with Session() as db:
        _seed(db)

    adapter = MockPlatformAdapter([_order_payload(order_no="PLATFORM-001")])
    monkeypatch.setattr(platform_sync_service, "get_platform_adapter", lambda platform: adapter)
    monkeypatch.setattr(platform_tasks, "SessionLocal", Session)
    first = platform_tasks.pull_shop_orders_task.apply(args=(1, 1), throw=True).result
    second = platform_tasks.pull_shop_orders_task.apply(args=(1, 1), throw=True).result
    assert first["created"] == 1
    assert second["updated"] == 1

    class WrongShopAdapter(MockPlatformAdapter):
        def pull_orders(self, shop, *, cursor=None):
            return [{**_order_payload(order_no="WRONG-SHOP"), "shop_id": 999}]

    monkeypatch.setattr(platform_sync_service, "get_platform_adapter", lambda platform: WrongShopAdapter())
    with Session() as db, pytest.raises(BizException):
        platform_sync_service.pull_shop_orders(db, shop_id=1, staff=db.get(Staff, 1))


def test_periodic_platform_task_queues_only_enabled_shops(monkeypatch):
    from types import SimpleNamespace
    from app.integrations import platform_tasks

    Session = _session_factory(pool=True)
    with Session() as db:
        _seed(db)
        db.add(Shop(id=2, platform_id=1, name="关闭同步的店铺", auto_sync_order=False))
        db.commit()
    queued = []
    monkeypatch.setattr(platform_tasks, "SessionLocal", Session)
    monkeypatch.setattr(platform_tasks, "get_settings", lambda: SimpleNamespace(platform_system_staff_id=1))
    monkeypatch.setattr(platform_tasks.pull_shop_orders_task, "delay", lambda shop_id, staff_id: queued.append((shop_id, staff_id)))
    result = platform_tasks.pull_enabled_shops_task.apply(throw=True).result
    assert result == {"queued": 1}
    assert queued == [(1, 1)]


def test_order_excel_import_export_and_template():
    Session = _session_factory()
    with Session() as db:
        _seed(db)
        staff = db.get(Staff, 1)
        template = order_service.ORDER_IMPORT_TEMPLATE
        wb = Workbook()
        ws = wb.active
        ws.title = template.sheet_title
        ws.append([column.header for column in template.columns])
        common = [
            "TB-EXCEL-001",
            "TAOBAO",
            "旗舰店",
            "2026-09-23 10:00:00",
            "2026-09-23 10:01:00",
            "PAID",
            "买家",
            "李四",
            "13800000001",
            "浙江省",
            "杭州市",
            "杭州市滨江区 1 号",
            "WH01",
            "经济快递",
            Decimal("5.00"),
            Decimal("0.00"),
            Decimal("0.00"),
            "",
            "",
        ]
        ws.append(
            common
            + [
                "P0001",
                "P0001-01",
                Decimal("1.000"),
                Decimal("10.00"),
                "否",
            ]
        )
        ws.append(
            common
            + [
                "P0002",
                "P0002-01",
                Decimal("2.000"),
                Decimal("8.00"),
                "是",
            ]
        )
        buf = BytesIO()
        wb.save(buf)

        result = order_service.import_orders(
            db, file_bytes=buf.getvalue(), staff=staff
        )
        assert result == {"created": 1, "updated": 0, "errors": []}
        order = order_service.get_order_or_404(db, 1)
        detail = order_service.order_dict(db, order)
        assert detail["order_amount"] == Decimal("31.00")
        assert detail["paid_amount"] == Decimal("31.00")
        assert detail["total_qty"] == Decimal("3.000")
        assert detail["items"][1]["is_gift"] is True

        exported = load_workbook(
            BytesIO(order_service.export_orders(db)),
            data_only=True,
        ).active
        assert exported.cell(row=1, column=1).value == "订单号"
        assert exported.cell(row=2, column=1).value == "TB-EXCEL-001"
        assert build_template(template).startswith(b"PK")


def _make_api():
    Session = _session_factory(pool=True)

    def override_get_db():
        db = Session()
        try:
            yield db
        finally:
            db.close()

    app = FastAPI()
    register_exception_handlers(app)
    app.include_router(order_router)
    app.dependency_overrides[get_db] = override_get_db
    return app, Session


def _auth(staff_id: int, name: str = "staff") -> dict:
    return {
        "Authorization": f"Bearer {create_access_token(staff_id, name, [])}"
    }


def test_api_order_permissions_audit_and_export(monkeypatch):
    monkeypatch.setattr(
        order_service,
        "generate_serial",
        lambda prefix, dt=None: "DD20260923-9001",
    )
    app, Session = _make_api()
    with Session() as db:
        _seed(db)
        _flow(db)
    client = TestClient(app)

    template = client.get("/orders/template", headers=_auth(1, "仓库"))
    assert template.status_code == 200
    assert template.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument"
    )

    denied_create = client.post(
        "/orders",
        json=jsonable_encoder(_order_payload(order_no=None)),
        headers=_auth(3, "售后"),
    )
    assert denied_create.status_code == 403

    created = client.post(
        "/orders",
        json=jsonable_encoder(_order_payload(order_no=None)),
        headers=_auth(1, "仓库"),
    )
    assert created.status_code == 200
    order_id = created.json()["data"]["id"]

    denied_audit = client.post(
        f"/orders/{order_id}/approve",
        json={"comment": "同意"},
        headers=_auth(1, "仓库"),
    )
    assert denied_audit.status_code == 403

    approved = client.post(
        f"/orders/{order_id}/approve",
        json={"comment": "同意"},
        headers=_auth(2, "财务"),
    )
    assert approved.status_code == 200
    assert approved.json()["data"]["status"] == "PENDING"

    logistics = client.post(
        f"/orders/{order_id}/logistics",
        json={"logistics_id": 2, "tracking_no": "SF999"},
        headers=_auth(1, "仓库"),
    )
    assert logistics.status_code == 200
    assert logistics.json()["data"]["tracking_no"] == "SF999"

    logs = client.get(f"/orders/{order_id}/logs", headers=_auth(1, "仓库"))
    assert logs.status_code == 200
    assert any(row["action"] == "SET_LOGISTICS" for row in logs.json()["data"])

    exported = client.get("/orders/export", headers=_auth(1, "仓库"))
    assert exported.status_code == 200
    assert exported.content.startswith(b"PK")
