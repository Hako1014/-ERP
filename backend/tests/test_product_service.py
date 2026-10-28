from decimal import Decimal

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, select
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

import app.models  # noqa: F401
from app.api.product import router as product_router
from app.core.errors import BizException
from app.core.handlers import register_exception_handlers
from app.core.security import create_access_token
from app.db.base import Base
from app.db.session import get_db
from app.models.basic import SpecType
from app.models.order import OrderItem
from app.models.org import Company, Staff
from app.models.product import ComboProduct, ComboProductItem, Product, ProductSku
from app.models.rbac import Role, StaffRole
from app.services import product_service
from app.services.excel import parse_data
from app.services.product_service import PRODUCT_TEMPLATE


def _session_factory(pool=None):
    kwargs = {"connect_args": {"check_same_thread": False}, "poolclass": StaticPool} if pool else {}
    engine = create_engine("sqlite:///:memory:", future=True, **kwargs)
    Base.metadata.create_all(engine)
    return sessionmaker(bind=engine, autoflush=False, expire_on_commit=False)


def _seed(db):
    db.add(Company(id=1, full_name="示例公司"))
    db.add(Staff(id=1, company_id=1, name="管理员", phone="13900000001"))
    db.commit()
    return db.get(Staff, 1)


def _base_data(**overrides) -> dict:
    data = {
        "name": "测试商品",
        "code": "P0001",
        "art_no": "ART-1",
        "online_price": "19.99",
        "tag_price": "29.99",
        "cost_price": "10.00",
        "warn_stock": 5,
    }
    data.update(overrides)
    return data


def test_create_product_with_sku():
    Session = _session_factory()
    with Session() as db:
        staff = _seed(db)
        product = product_service.create_product(
            db,
            data=_base_data(),
            staff=staff,
            skus=[{"spec_code": "P0001-RED", "spec_name": "红色", "sale_price": "19.90"}],
        )
        detail = product_service.product_dict(db, product, with_skus=True)
        assert detail["code"] == "P0001"
        assert detail["online_price"] == Decimal("19.99")
        assert detail["listed"] is False
        assert len(detail["skus"]) == 1
        assert detail["skus"][0]["spec_code"] == "P0001-RED"


def test_duplicate_code_rejected():
    Session = _session_factory()
    with Session() as db:
        staff = _seed(db)
        product_service.create_product(db, data=_base_data(), staff=staff)
        with pytest.raises(BizException) as exc:
            product_service.create_product(
                db, data=_base_data(name="另一个商品"), staff=staff
            )
        assert "商品编码已存在" in exc.value.message


def test_create_product_requires_name_and_code():
    Session = _session_factory()
    with Session() as db:
        staff = _seed(db)
        with pytest.raises(BizException) as exc:
            product_service.create_product(db, data={"name": "缺编码"}, staff=staff)
        assert "必填" in exc.value.message


def test_float_money_rejected():
    Session = _session_factory()
    with Session() as db:
        staff = _seed(db)
        with pytest.raises(BizException) as exc:
            product_service.create_product(
                db, data=_base_data(online_price=19.99), staff=staff
            )
        assert "金额" in exc.value.message


def test_multi_measure_requires_unit_group():
    Session = _session_factory()
    with Session() as db:
        staff = _seed(db)
        with pytest.raises(BizException) as exc:
            product_service.create_product(
                db, data=_base_data(measure_type="MULTI"), staff=staff
            )
        assert "计量单位组" in exc.value.message

        unit = product_service.create_unit(db, name="件")
        group = product_service.create_unit_group(
            db, name="基本组", main_unit_id=unit.id
        )
        product = product_service.create_product(
            db,
            data=_base_data(measure_type="MULTI", unit_group_id=group.id),
            staff=staff,
        )
        assert product.main_unit_id == unit.id


def test_listing_and_disable():
    Session = _session_factory()
    with Session() as db:
        staff = _seed(db)
        product = product_service.create_product(db, data=_base_data(), staff=staff)

        assert product_service.set_listing(
            db, product_ids=[product.id], listed=True, staff=staff
        ) == 1
        assert product.listed_at is not None

        product_service.set_disabled(
            db, product_id=product.id, disabled=True, staff=staff
        )
        assert product.listed_at is None
        assert product.is_disabled is True

        with pytest.raises(BizException) as exc:
            product_service.set_listing(
                db, product_ids=[product.id], listed=True, staff=staff
            )
        assert "已禁用" in exc.value.message


def test_list_products_filters_and_paging():
    Session = _session_factory()
    with Session() as db:
        staff = _seed(db)
        product_service.create_product(db, data=_base_data(code="P0001"), staff=staff)
        product_service.create_product(
            db, data=_base_data(code="P0002", name="苹果"), staff=staff
        )
        result = product_service.list_products(db, keyword="苹果")
        assert result["total"] == 1
        assert result["items"][0]["code"] == "P0002"

        result = product_service.list_products(db, page=2, page_size=1)
        assert result["total"] == 2
        assert len(result["items"]) == 1


def test_delete_product_removes_skus():
    Session = _session_factory()
    with Session() as db:
        staff = _seed(db)
        product = product_service.create_product(
            db,
            data=_base_data(),
            staff=staff,
            skus=[{"spec_code": "P0001-A", "spec_name": "A"}],
        )
        sku_id = db.scalar(select(ProductSku.id).where(ProductSku.product_id == product.id))
        result = product_service.delete_products(db, product_ids=[product.id], staff=staff)
        assert result["deleted"] == 1
        assert db.get(ProductSku, sku_id) is None
        assert db.get(Product, product.id) is None


def test_delete_product_blocked_when_referenced():
    Session = _session_factory()
    with Session() as db:
        staff = _seed(db)
        product = product_service.create_product(db, data=_base_data(), staff=staff)
        db.add(
            OrderItem(
                order_id=1,
                product_id=product.id,
                qty=Decimal("1"),
                price=Decimal("10.00"),
            )
        )
        db.commit()
        with pytest.raises(BizException) as exc:
            product_service.delete_products(db, product_ids=[product.id], staff=staff)
        assert "不能删除" in exc.value.message


def test_sku_code_unique_and_update():
    Session = _session_factory()
    with Session() as db:
        staff = _seed(db)
        product = product_service.create_product(db, data=_base_data(), staff=staff)
        sku = product_service.create_sku(
            db, product_id=product.id, data={"spec_code": "SKU-1"}
        )
        with pytest.raises(BizException) as exc:
            product_service.create_sku(
                db, product_id=product.id, data={"spec_code": "SKU-1"}
            )
        assert "规格编码已存在" in exc.value.message

        updated = product_service.update_sku(
            db, sku_id=sku.id, data={"sale_price": "88.88", "is_active": False}
        )
        assert updated.sale_price == Decimal("88.88")
        assert updated.is_active is False


def test_export_and_import_roundtrip():
    Session = _session_factory()
    with Session() as db:
        staff = _seed(db)
        category = product_service.create_category(db, name="数码")
        brand = product_service.create_brand(db, name="示例品牌")
        unit = product_service.create_unit(db, name="件")
        product_service.create_product(
            db,
            data=_base_data(
                category_id=category.id, brand_id=brand.id, main_unit_id=unit.id
            ),
            staff=staff,
        )

        content = product_service.export_products(db)
        parsed = parse_data(content, PRODUCT_TEMPLATE)
        assert parsed.errors == []
        assert parsed.rows[0]["code"] == "P0001"
        assert parsed.rows[0]["category_name"] == "数码"
        assert parsed.rows[0]["unit_name"] == "件"

        parsed.rows[0]["code"] = "P0002"
        parsed.rows[0]["name"] = "导入商品"
        from io import BytesIO

        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = PRODUCT_TEMPLATE.sheet_title
        ws.append([c.header for c in PRODUCT_TEMPLATE.columns])
        ws.append([parsed.rows[0].get(c.field) for c in PRODUCT_TEMPLATE.columns])
        buf = BytesIO()
        wb.save(buf)

        report = product_service.import_products(db, file_bytes=buf.getvalue(), staff=staff)
        assert report["created"] == 1
        assert report["errors"] == []
        assert db.scalar(select(Product).where(Product.code == "P0002")) is not None


def test_import_rejects_unknown_category_without_writing():
    Session = _session_factory()
    with Session() as db:
        staff = _seed(db)
        from io import BytesIO

        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = PRODUCT_TEMPLATE.sheet_title
        ws.append([c.header for c in PRODUCT_TEMPLATE.columns])
        ws.append(["P1001", "导入商品", None, "不存在的分类"])
        buf = BytesIO()
        wb.save(buf)

        report = product_service.import_products(db, file_bytes=buf.getvalue(), staff=staff)
        assert report["created"] == 0
        assert report["errors"][0]["message"].startswith("分类不存在")
        assert db.scalar(select(Product).where(Product.code == "P1001")) is None


def test_import_rejects_existing_code():
    Session = _session_factory()
    with Session() as db:
        staff = _seed(db)
        product_service.create_product(db, data=_base_data(), staff=staff)
        from io import BytesIO

        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = PRODUCT_TEMPLATE.sheet_title
        ws.append([c.header for c in PRODUCT_TEMPLATE.columns])
        ws.append(["P0001", "重复商品"])
        buf = BytesIO()
        wb.save(buf)

        report = product_service.import_products(db, file_bytes=buf.getvalue(), staff=staff)
        assert report["created"] == 0
        assert "已存在" in report["errors"][0]["message"]


def test_settings_crud_and_guards():
    Session = _session_factory()
    with Session() as db:
        staff = _seed(db)
        parent = product_service.create_category(db, name="服装")
        child = product_service.create_category(db, name="男装", parent_id=parent.id)
        with pytest.raises(BizException):
            product_service.delete_category(db, category_id=parent.id)

        brand = product_service.create_brand(db, name="品牌A")
        unit = product_service.create_unit(db, name="件")
        group = product_service.create_unit_group(
            db,
            name="组合单位",
            main_unit_id=unit.id,
            items=[{"unit_id": unit.id, "conversion_rate": "1"}],
        )
        assert product_service.list_unit_groups(db)[0]["items"][0]["unit_id"] == unit.id

        with pytest.raises(BizException) as exc:
            product_service.create_unit_group(
                db, name="非法单位组", main_unit_id=unit.id,
                items=[{"unit_id": unit.id, "conversion_rate": "0"}],
            )
        assert "换算率" in exc.value.message

        spec_type = product_service.create_spec_type(
            db, name="颜色", category_id=parent.id
        )
        value = product_service.create_spec_value(
            db, spec_type_id=spec_type.id, name="红色", code="RED"
        )
        assert product_service.list_spec_values(db, spec_type.id)[0]["code"] == "RED"

        product = product_service.create_product(
            db, data=_base_data(category_id=child.id, brand_id=brand.id), staff=staff
        )
        with pytest.raises(BizException):
            product_service.delete_brand(db, brand_id=brand.id)
        with pytest.raises(BizException):
            product_service.delete_category(db, category_id=child.id)

        product_service.delete_spec_value(db, spec_value_id=value.id)
        product_service.delete_spec_type(db, spec_type_id=spec_type.id)
        assert db.get(SpecType, spec_type.id) is None
        assert db.get(Product, product.id) is not None


def _make_api():
    Session = _session_factory(pool=True)

    def override_get_db():
        db = Session()
        try:
            yield db
        finally:
            db.close()

    app = FastAPI()
    register_exception_handlers(app)
    app.include_router(product_router)
    app.dependency_overrides[get_db] = override_get_db
    return app, Session


def _seed_roles(Session):
    with Session() as db:
        _seed(db)
        db.add_all(
            [
                Role(id=1, company_id=1, code="ADMIN", name="管理员"),
                Role(id=2, company_id=1, code="PURCHASER", name="采购员"),
            ]
        )
        db.add(StaffRole(staff_id=1, role_id=2))
        db.commit()


def _auth(staff_id: int, name: str = "staff") -> dict:
    return {"Authorization": f"Bearer {create_access_token(staff_id, name, [])}"}


def test_api_read_allowed_but_write_denied_for_purchaser():
    app, Session = _make_api()
    _seed_roles(Session)
    client = TestClient(app)

    listed = client.get("/products", headers=_auth(1))
    assert listed.status_code == 200
    assert listed.json()["data"]["total"] == 0

    denied = client.post(
        "/products",
        json={"name": "商品", "code": "P9001"},
        headers=_auth(1),
    )
    assert denied.status_code == 403


def test_api_admin_can_create_and_list_product():
    app, Session = _make_api()
    with Session() as db:
        db.add_all(
            [
                Company(id=1, full_name="示例公司"),
                Staff(id=9, company_id=1, name="管理员", phone="13900000009"),
                Role(id=1, company_id=1, code="ADMIN", name="管理员"),
            ]
        )
        db.flush()
        db.add(StaffRole(staff_id=9, role_id=1))
        db.commit()

    client = TestClient(app)
    created = client.post(
        "/products",
        json={"name": "商品A", "code": "P9002", "online_price": "12.30"},
        headers=_auth(9, "管理员"),
    )
    assert created.status_code == 200
    assert created.json()["data"]["online_price"] == "12.30"

    detail = client.get("/products", headers=_auth(9, "管理员"))
    assert detail.json()["data"]["total"] == 1


# --------------------------------------------------------------------------- #
# 组合商品
# --------------------------------------------------------------------------- #


def _three_products(db, staff) -> tuple[Product, Product, Product]:
    main = product_service.create_product(
        db, data=_base_data(code="C-MAIN", name="主机"), staff=staff
    )
    acc1 = product_service.create_product(
        db, data=_base_data(code="C-ACC1", name="电源"), staff=staff
    )
    acc2 = product_service.create_product(
        db, data=_base_data(code="C-ACC2", name="数据线"), staff=staff
    )
    return main, acc1, acc2


def test_create_combo_and_expand():
    Session = _session_factory()
    with Session() as db:
        staff = _seed(db)
        main, acc1, acc2 = _three_products(db, staff)
        combo = product_service.create_combo(
            db,
            data={
                "combo_no": "CB-001",
                "name": "主机套装",
                "main_product_id": main.id,
            },
            items=[
                {"product_id": acc1.id, "qty": "1"},
                {"product_id": acc2.id, "qty": "2"},
            ],
            staff=staff,
        )
        detail = product_service.combo_dict(db, combo)
        assert detail["combo_no"] == "CB-001"
        assert len(detail["items"]) == 2

        parts = product_service.expand_combo(db, product_id=main.id, qty="2")
        assert parts[0]["kind"] == "MAIN"
        assert parts[0]["qty"] == Decimal("2")
        accessory_qty = {p["product_id"]: p["qty"] for p in parts if p["kind"] == "ACCESSORY"}
        assert accessory_qty[acc1.id] == Decimal("2")
        assert accessory_qty[acc2.id] == Decimal("4")


def test_expand_normal_product_returns_single():
    Session = _session_factory()
    with Session() as db:
        staff = _seed(db)
        product = product_service.create_product(db, data=_base_data(), staff=staff)
        parts = product_service.expand_combo(db, product_id=product.id, qty="3")
        assert parts == [
            {
                "combo_id": None,
                "kind": "SINGLE",
                "product_id": product.id,
                "sku_id": None,
                "qty": Decimal("3"),
            }
        ]


def test_combo_rejects_nested_and_invalid_items():
    Session = _session_factory()
    with Session() as db:
        staff = _seed(db)
        main, acc1, acc2 = _three_products(db, staff)
        with pytest.raises(BizException) as exc:
            product_service.create_combo(
                db,
                data={"combo_no": "CB-002", "name": "空组合", "main_product_id": main.id},
                items=[],
                staff=staff,
            )
        assert "至少需要一个配件" in exc.value.message

        with pytest.raises(BizException) as exc:
            product_service.create_combo(
                db,
                data={"combo_no": "CB-003", "name": "数量非法", "main_product_id": main.id},
                items=[{"product_id": acc1.id, "qty": "0"}],
                staff=staff,
            )
        assert "数量必须大于 0" in exc.value.message

        product_service.create_combo(
            db,
            data={"combo_no": "CB-004", "name": "内层组合", "main_product_id": acc1.id},
            items=[{"product_id": acc2.id, "qty": "1"}],
            staff=staff,
        )
        with pytest.raises(BizException) as exc:
            product_service.create_combo(
                db,
                data={"combo_no": "CB-005", "name": "嵌套组合", "main_product_id": main.id},
                items=[{"product_id": acc1.id, "qty": "1"}],
                staff=staff,
            )
        assert "嵌套组合" in exc.value.message


def test_combo_no_unique_and_update_items():
    Session = _session_factory()
    with Session() as db:
        staff = _seed(db)
        main, acc1, acc2 = _three_products(db, staff)
        combo = product_service.create_combo(
            db,
            data={"combo_no": "CB-100", "name": "套装A", "main_product_id": main.id},
            items=[{"product_id": acc1.id, "qty": "1"}],
            staff=staff,
        )
        with pytest.raises(BizException) as exc:
            product_service.create_combo(
                db,
                data={"combo_no": "CB-100", "name": "套装B", "main_product_id": main.id},
                items=[{"product_id": acc2.id, "qty": "1"}],
                staff=staff,
            )
        assert "组合编码已存在" in exc.value.message

        product_service.update_combo(
            db,
            combo_id=combo.id,
            data={"name": "套装A-改"},
            items=[{"product_id": acc2.id, "qty": "3"}],
            staff=staff,
        )
        detail = product_service.combo_dict(db, combo)
        assert detail["name"] == "套装A-改"
        assert [item["product_id"] for item in detail["items"]] == [acc2.id]
        assert detail["items"][0]["qty"] == Decimal("3")


def test_disable_combo_stops_expansion():
    Session = _session_factory()
    with Session() as db:
        staff = _seed(db)
        main, acc1, _ = _three_products(db, staff)
        combo = product_service.create_combo(
            db,
            data={"combo_no": "CB-200", "name": "套装C", "main_product_id": main.id},
            items=[{"product_id": acc1.id, "qty": "1"}],
            staff=staff,
        )
        product_service.set_combo_enabled(db, combo_id=combo.id, enabled=False)
        parts = product_service.expand_combo(db, product_id=main.id, qty="1")
        assert parts[0]["kind"] == "SINGLE"


def test_delete_combo_removes_items():
    Session = _session_factory()
    with Session() as db:
        staff = _seed(db)
        main, acc1, _ = _three_products(db, staff)
        combo = product_service.create_combo(
            db,
            data={"combo_no": "CB-300", "name": "套装D", "main_product_id": main.id},
            items=[{"product_id": acc1.id, "qty": "1"}],
            staff=staff,
        )
        product_service.delete_combo(db, combo_id=combo.id)
        assert db.get(ComboProduct, combo.id) is None
        assert (
            db.scalar(
                select(ComboProductItem).where(ComboProductItem.combo_id == combo.id)
            )
            is None
        )


def test_api_combo_create_and_expand():
    app, Session = _make_api()
    with Session() as db:
        db.add_all(
            [
                Company(id=1, full_name="示例公司"),
                Staff(id=9, company_id=1, name="管理员", phone="13900000009"),
                Role(id=1, company_id=1, code="ADMIN", name="管理员"),
            ]
        )
        db.flush()
        db.add(StaffRole(staff_id=9, role_id=1))
        db.commit()
        staff = db.get(Staff, 9)
        main, acc1, _ = _three_products(db, staff)

    client = TestClient(app)
    created = client.post(
        "/combo-products",
        json={
            "combo_no": "CB-API-1",
            "name": "接口套装",
            "main_product_id": main.id,
            "items": [{"product_id": acc1.id, "qty": "2"}],
        },
        headers=_auth(9, "管理员"),
    )
    assert created.status_code == 200
    combo_id = created.json()["data"]["id"]

    expanded = client.get(
        f"/combo-products/{combo_id}/expand?qty=3", headers=_auth(9, "管理员")
    )
    assert expanded.status_code == 200
    qty_map = {p["kind"]: Decimal(str(p["qty"])) for p in expanded.json()["data"]}
    assert qty_map["MAIN"] == Decimal("3")
    assert qty_map["ACCESSORY"] == Decimal("6")
