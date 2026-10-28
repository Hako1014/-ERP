from datetime import datetime
from decimal import Decimal
from io import BytesIO

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from PIL import Image
from sqlalchemy import create_engine, event, select
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

import app.models  # noqa: F401
import app.api.purchase_order as purchase_order_api
from app.api.purchase_order import router as purchase_order_router
from app.core.data_scope import initialize_scope_policy
from app.core.errors import BizException
from app.core.handlers import register_exception_handlers
from app.core.security import create_access_token
from app.db.base import Base, CompanyOwnedBase
from app.db.session import get_db
from app.models.basic import Brand, Unit
from app.models.org import Company, Staff
from app.models.party import Supplier, SupplierPrice
from app.models.product import Product, ProductSku
from app.models.purchase import PurchasePlan, PurchasePlanItem
from app.models.rbac import Role, StaffRole
from app.models.warehouse import Warehouse
from app.services import (
    approval_service,
    purchase_order_service,
)
from app.services.storage import LocalStorage


def _session_factory(pool: bool = False):
    kwargs = (
        {"connect_args": {"check_same_thread": False}, "poolclass": StaticPool}
        if pool
        else {}
    )
    engine = create_engine("sqlite:///:memory:", future=True, **kwargs)
    Base.metadata.create_all(engine)
    session_factory = sessionmaker(bind=engine, autoflush=False, expire_on_commit=False)

    def backfill_legacy_fixture_rows(db, _flush_context, _instances):
        for row in db.new:
            if isinstance(row, CompanyOwnedBase) and row.company_id is None:
                row.company_id = 1

    event.listen(session_factory.class_, "before_flush", backfill_legacy_fixture_rows)
    return session_factory


def _seed(db) -> dict:
    db.add(Company(id=1, full_name="示例公司"))
    db.add_all(
        [
            Staff(id=1, company_id=1, name="采购员", phone="13900000001"),
            Staff(id=2, company_id=1, name="计划审核员", phone="13900000002"),
            Staff(id=3, company_id=1, name="财务", phone="13900000003"),
            Staff(id=9, company_id=1, name="管理员", phone="13900000009"),
        ]
    )
    db.add(Supplier(id=1, code="S0001", name="示例供应商", is_active=True))
    db.add_all(
        [
            Brand(id=1, name="示例品牌"),
            Unit(id=1, name="件", code="PCS"),
        ]
    )
    db.add_all(
        [
            Product(
                id=1,
                name="测试商品A",
                code="P0001",
                art_no="ART-1",
                brand_id=1,
                main_unit_id=1,
                default_supplier_id=1,
                cost_price=Decimal("8.00"),
            ),
            Product(
                id=2,
                name="测试商品B",
                code="P0002",
                art_no="ART-2",
                brand_id=1,
                main_unit_id=1,
                default_supplier_id=1,
                cost_price=Decimal("6.00"),
            ),
        ]
    )
    db.add_all(
        [
            ProductSku(
                id=1,
                product_id=1,
                spec_name="红色",
                spec_code="P0001-RED",
                cost_price=Decimal("7.50"),
                is_active=True,
            ),
            ProductSku(
                id=2,
                product_id=2,
                spec_name="蓝色",
                spec_code="P0002-BLUE",
                cost_price=Decimal("5.50"),
                is_active=True,
            ),
        ]
    )
    db.add_all(
        [
            SupplierPrice(
                supplier_id=1,
                product_id=1,
                sku_id=1,
                price=Decimal("10.00"),
                status="ACTIVE",
            ),
            SupplierPrice(
                supplier_id=1,
                product_id=2,
                sku_id=2,
                price=Decimal("8.00"),
                status="ACTIVE",
            ),
        ]
    )
    db.add_all(
        [
            Warehouse(
                id=1,
                code="WH01",
                name="主仓",
                is_enabled=True,
                is_virtual=False,
                lock_type="NONE",
            ),
            Warehouse(
                id=2,
                code="WH02",
                name="其他仓",
                is_enabled=True,
                is_virtual=False,
                lock_type="NONE",
            ),
        ]
    )
    db.add(PurchasePlan(id=1, plan_no="CGJH001", supplier_id=1))
    db.add_all(
        [
            PurchasePlanItem(
                id=1,
                plan_id=1,
                product_id=1,
                sku_id=1,
                planned_qty=Decimal("10.000"),
                purchased_qty=Decimal("0.000"),
                purchase_price=Decimal("10.00"),
                amount=Decimal("100.00"),
            ),
            PurchasePlanItem(
                id=2,
                plan_id=1,
                product_id=2,
                sku_id=2,
                planned_qty=Decimal("5.000"),
                purchased_qty=Decimal("0.000"),
                purchase_price=Decimal("8.00"),
                amount=Decimal("40.00"),
            ),
        ]
    )
    db.flush()
    plan = db.get(PurchasePlan, 1)
    plan.status = "IN_PROGRESS"
    plan.audit_status = "APPROVED"
    plan.total_qty = Decimal("15.000")
    plan.total_amount = Decimal("140.00")
    db.commit()
    initialize_scope_policy(db, db.get(Staff, 1))
    return {"supplier_id": 1, "warehouse_id": 1}


def _approved_plan(db) -> None:
    plan = db.get(PurchasePlan, 1)
    plan.status = "IN_PROGRESS"
    plan.audit_status = "APPROVED"
    db.flush()


def test_plan_to_order_approval_arrival_and_warehouse_lock(monkeypatch):
    serials = iter(["CGDD20260923-0001", "CGDD20260923-0002"])
    monkeypatch.setattr(
        purchase_order_service,
        "generate_serial",
        lambda prefix, dt=None: next(serials),
    )
    Session = _session_factory()
    with Session() as db:
        _seed(db)
        purchaser = db.get(Staff, 1)
        finance = db.get(Staff, 3)
        approval_service.create_flow(
            db,
            name="采购订单审批",
            doc_type="purchase_order",
            approvers=[finance.id],
            actor=purchaser,
        )
        order = purchase_order_service.create_order_from_plan(
            db,
            data={
                "plan_id": 1,
                "warehouse_id": 1,
                "items": [{"plan_item_id": 1, "qty": "6.000"}],
            },
            staff=purchaser,
        )
        assert order.order_no == "CGDD20260923-0001"
        assert order.status == "PENDING"
        assert order.audit_status == "PENDING"
        assert order.total_qty == Decimal("6.000")
        assert order.total_amount == Decimal("60.00")
        detail = purchase_order_service.order_dict(db, order)
        assert detail["items"][0]["in_transit_qty"] == Decimal("6.000")
        assert detail["items"][0]["difference_qty"] == Decimal("6.000")
        assert db.get(PurchasePlanItem, 1).purchased_qty == Decimal("6.000")

        order = purchase_order_service.submit_order(
            db, order_id=order.id, staff=purchaser
        )
        assert purchase_order_service.order_dict(
            db, order, with_approval=True
        )["approval"]["state"]["current_step"] == 1
        order = purchase_order_service.approve_order(
            db, order_id=order.id, staff=finance, comment="同意"
        )
        assert order.status == "APPROVED"
        assert order.audit_status == "APPROVED"
        assert order.auditor_id == finance.id

        order = purchase_order_service.record_arrival(
            db,
            order_id=order.id,
            warehouse_id=1,
            items=[{"order_item_id": 1, "qty": "2.000"}],
            staff=finance,
        )
        detail = purchase_order_service.order_dict(db, order)
        assert order.status == "RECEIVING"
        assert order.arrived_qty == Decimal("2.000")
        assert detail["difference_qty"] == Decimal("4.000")
        assert detail["items"][0]["arrived_qty"] == Decimal("2.000")
        assert detail["items"][0]["difference_qty"] == Decimal("4.000")

        with pytest.raises(BizException) as mismatch:
            purchase_order_service.record_arrival(
                db,
                order_id=order.id,
                warehouse_id=2,
                items=[{"order_item_id": 1, "qty": "1.000"}],
                staff=finance,
            )
        assert "仓库" in mismatch.value.message

        order = purchase_order_service.record_arrival(
            db,
            order_id=order.id,
            warehouse_id=1,
            items=[{"order_item_id": 1, "qty": "4.000"}],
            staff=finance,
        )
        assert order.status == "RECEIVED"
        assert order.arrived_qty == Decimal("6.000")

        with pytest.raises(BizException) as unaudit_exc:
            purchase_order_service.unaudit_order(
                db, order_id=order.id, staff=purchaser, reason="调整"
            )
        assert "到货" in unaudit_exc.value.message


def test_update_delete_and_plan_purchased_qty_delta(monkeypatch):
    monkeypatch.setattr(
        purchase_order_service,
        "generate_serial",
        lambda prefix, dt=None: "CGDD20260923-1001",
    )
    Session = _session_factory()
    with Session() as db:
        _seed(db)
        purchaser = db.get(Staff, 1)
        order = purchase_order_service.create_order_from_plan(
            db,
            data={
                "plan_id": 1,
                "warehouse_id": 1,
                "items": [{"plan_item_id": 1, "qty": "8.000"}],
            },
            staff=purchaser,
        )
        assert db.get(PurchasePlanItem, 1).purchased_qty == Decimal("8.000")

        order = purchase_order_service.update_order(
            db,
            order_id=order.id,
            data={
                "items": [
                    {
                        "product_id": 1,
                        "sku_id": 1,
                        "qty": "9.000",
                        "purchase_price": "10.00",
                    }
                ]
            },
            staff=purchaser,
        )
        assert order.total_qty == Decimal("9.000")
        assert db.get(PurchasePlanItem, 1).purchased_qty == Decimal("9.000")

        with pytest.raises(BizException) as over_exc:
            purchase_order_service.update_order(
                db,
                order_id=order.id,
                data={
                    "items": [
                        {
                            "product_id": 1,
                            "sku_id": 1,
                            "qty": "11.000",
                            "purchase_price": "10.00",
                        }
                    ]
                },
                staff=purchaser,
            )
        assert "超过计划剩余" in over_exc.value.message
        assert db.get(PurchasePlanItem, 1).purchased_qty == Decimal("9.000")

        purchase_order_service.delete_order(
            db, order_id=order.id, staff=purchaser
        )
        assert db.get(PurchasePlanItem, 1).purchased_qty == Decimal("0.000")


def test_import_export_template_and_pdf(monkeypatch):
    monkeypatch.setattr(
        purchase_order_service,
        "generate_serial",
        lambda prefix, dt=None: "CGDD20260923-2001",
    )
    Session = _session_factory()
    with Session() as db:
        _seed(db)
        staff = db.get(Staff, 1)
        template = purchase_order_service.PURCHASE_ORDER_TEMPLATE
        wb = Workbook()
        ws = wb.active
        ws.title = template.sheet_title
        ws.append([column.header for column in template.columns])
        ws.append(
            [
                "S0001",
                "WH01",
                "",
                "Excel导入采购订单",
                "P0001",
                "P0001-RED",
                Decimal("10.00"),
                Decimal("3.000"),
                "/uploads/a.jpg|/uploads/b.jpg",
            ]
        )
        ws.append(
            [
                "S0001",
                "WH01",
                "",
                "Excel导入采购订单",
                "P0002",
                "P0002-BLUE",
                Decimal("8.00"),
                Decimal("2.000"),
                "/uploads/a.jpg|/uploads/b.jpg",
            ]
        )
        buf = BytesIO()
        wb.save(buf)

        result = purchase_order_service.import_orders(
            db, file_bytes=buf.getvalue(), staff=staff
        )
        assert result["created"] == 1
        assert result["item_count"] == 2
        order = purchase_order_service.get_order_or_404(db, result["order_id"])
        detail = purchase_order_service.order_dict(db, order)
        assert detail["total_qty"] == Decimal("5.000")
        assert detail["attachments"] == [
            "/uploads/a.jpg",
            "/uploads/b.jpg",
        ]

        export = purchase_order_service.export_orders(db)
        exported = load_workbook(BytesIO(export), data_only=True).active
        assert exported.cell(row=1, column=1).value == "采购订单号"
        assert exported.max_row == 2

        pdf = purchase_order_service.build_order_pdf(db, order.id)
        assert pdf.startswith(b"%PDF")


def _make_api(monkeypatch, tmp_path):
    Session = _session_factory(pool=True)

    def override_get_db():
        db = Session()
        try:
            yield db
        finally:
            db.close()

    app = FastAPI()
    register_exception_handlers(app)
    app.include_router(purchase_order_router)
    app.dependency_overrides[get_db] = override_get_db
    monkeypatch.setattr(
        purchase_order_api,
        "get_storage",
        lambda: LocalStorage(tmp_path),
    )
    monkeypatch.setattr(
        purchase_order_service,
        "generate_serial",
        lambda prefix, dt=None: "CGDD20260923-3001",
    )
    return app, Session


def _auth(staff_id: int, name: str = "staff") -> dict:
    return {"Authorization": f"Bearer {create_access_token(staff_id, name, [])}"}


def _seed_api(Session):
    with Session() as db:
        _seed(db)
        db.add_all(
            [
                Role(id=1, company_id=1, code="ADMIN", name="管理员"),
                Role(id=2, company_id=1, code="PURCHASER", name="采购员"),
                Role(id=3, company_id=1, code="FINANCE", name="财务"),
            ]
        )
        db.flush()
        db.add_all(
            [
                StaffRole(staff_id=1, role_id=2),
                StaffRole(staff_id=3, role_id=3),
                StaffRole(staff_id=9, role_id=1),
            ]
        )
        db.commit()
        approval_service.create_flow(
            db,
            name="采购订单接口审批",
            doc_type="purchase_order",
            approvers=[3],
            actor=db.get(Staff, 1),
        )


def test_api_purchase_order_permissions_attachment_and_download(monkeypatch, tmp_path):
    app, Session = _make_api(monkeypatch, tmp_path)
    _seed_api(Session)
    client = TestClient(app)

    created = client.post(
        "/purchase-orders",
        json={
            "supplier_id": 1,
            "warehouse_id": 1,
            "description": "接口创建",
            "items": [
                {
                    "product_id": 1,
                    "sku_id": 1,
                    "qty": "2.000",
                    "purchase_price": "10.00",
                }
            ],
        },
        headers=_auth(1, "采购员"),
    )
    assert created.status_code == 200
    order_id = created.json()["data"]["id"]

    denied = client.post(
        f"/purchase-orders/{order_id}/approve",
        json={"comment": "同意"},
        headers=_auth(1, "采购员"),
    )
    assert denied.status_code == 403

    image = BytesIO()
    Image.new("RGB", (8, 8), "red").save(image, format="PNG")
    uploaded = client.post(
        f"/purchase-orders/{order_id}/attachments",
        files={"files": ("supplier.png", image.getvalue(), "image/png")},
        headers=_auth(1, "采购员"),
    )
    assert uploaded.status_code == 200
    assert len(uploaded.json()["data"]["attachments"]) == 1

    submitted = client.post(
        f"/purchase-orders/{order_id}/submit",
        headers=_auth(1, "采购员"),
    )
    assert submitted.status_code == 200

    approved = client.post(
        f"/purchase-orders/{order_id}/approve",
        json={"comment": "同意"},
        headers=_auth(3, "财务"),
    )
    assert approved.status_code == 200
    assert approved.json()["data"]["status"] == "APPROVED"

    exported = client.get("/purchase-orders/export", headers=_auth(1))
    assert exported.status_code == 200
    assert exported.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument"
    )

    pdf = client.get(
        f"/purchase-orders/{order_id}/pdf", headers=_auth(1)
    )
    assert pdf.status_code == 200
    assert pdf.content.startswith(b"%PDF")
