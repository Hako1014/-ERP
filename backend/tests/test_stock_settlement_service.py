from datetime import datetime
from decimal import Decimal
from io import BytesIO

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

import app.models  # noqa: F401
from app.api.stock_settlement import router as stock_settlement_router
from app.core.errors import BizException
from app.core.handlers import register_exception_handlers
from app.core.security import create_access_token
from app.db.base import Base
from app.db.session import get_db
from app.models.basic import Unit
from app.models.org import Company, Staff
from app.models.platform import Platform, Shop
from app.models.product import Product, ProductSku
from app.models.rbac import Role, StaffRole
from app.models.stock import (
    ShopStock,
    StockFlow,
    StockOut,
    StockOutItem,
    StockSettlement,
)
from app.models.warehouse import Location, StockLocation, Warehouse
from app.services import stock_settlement_service
from app.services.excel import build_template

XLSX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)


def _session_factory(pool: bool = False):
    kwargs = (
        {"connect_args": {"check_same_thread": False}, "poolclass": StaticPool}
        if pool
        else {}
    )
    engine = create_engine("sqlite:///:memory:", future=True, **kwargs)
    Base.metadata.create_all(engine)
    return sessionmaker(bind=engine, autoflush=False, expire_on_commit=False)


def _seed(db) -> None:
    db.add(Company(id=1, full_name="示例公司"))
    db.add_all(
        [
            Staff(id=1, company_id=1, name="仓库", phone="13900000001"),
            Staff(id=2, company_id=1, name="管理员", phone="13900000002"),
            Staff(id=3, company_id=1, name="质检", phone="13900000003"),
        ]
    )
    db.add(Unit(id=1, code="PCS", name="件"))
    db.add_all(
        [
            Product(
                id=1,
                name="主商品",
                code="P0001",
                main_unit_id=1,
                cost_price=Decimal("10.00"),
            ),
            Product(
                id=2,
                name="配件商品",
                code="P0002",
                main_unit_id=1,
                cost_price=Decimal("5.00"),
            ),
        ]
    )
    db.add_all(
        [
            ProductSku(
                id=1,
                product_id=1,
                spec_name="默认",
                spec_code="P0001-01",
                barcode="690000000001",
                cost_price=Decimal("10.00"),
                is_active=True,
            ),
            ProductSku(
                id=2,
                product_id=2,
                spec_name="默认",
                spec_code="P0002-01",
                barcode="690000000002",
                cost_price=Decimal("5.00"),
                is_active=True,
            ),
        ]
    )
    db.add_all(
        [
            Warehouse(
                id=1,
                code="WH01",
                name="主仓",
                warehouse_type="FINISHED",
                is_enabled=True,
                enable_location=True,
                lock_type="NONE",
            ),
            Warehouse(
                id=2,
                code="WH02",
                name="次品仓",
                warehouse_type="DEFECTIVE",
                is_enabled=True,
                enable_location=True,
                lock_type="NONE",
            ),
        ]
    )
    db.add_all(
        [
            Location(id=1, warehouse_id=1, code="WH01-A01", is_enabled=True),
            Location(id=2, warehouse_id=1, code="WH01-A02", is_enabled=True),
            Location(id=3, warehouse_id=2, code="WH02-A01", is_enabled=True),
        ]
    )
    db.add_all(
        [
            StockLocation(
                id=1,
                location_id=1,
                product_sku_id=1,
                qty=Decimal("10.000"),
                status="NORMAL",
            ),
            StockLocation(
                id=2,
                location_id=2,
                product_sku_id=1,
                qty=Decimal("5.000"),
                status="NORMAL",
            ),
            StockLocation(
                id=3,
                location_id=1,
                product_sku_id=2,
                qty=Decimal("3.000"),
                status="NORMAL",
            ),
            StockLocation(
                id=4,
                location_id=3,
                product_sku_id=1,
                qty=Decimal("2.000"),
                status="DEFECTIVE",
            ),
        ]
    )
    db.add_all(
        [
            Platform(id=1, name="淘宝", code="TAOBAO"),
            Platform(id=2, name="京东", code="JD"),
        ]
    )
    db.add_all(
        [
            Shop(id=1, platform_id=1, name="旗舰店", auto_sync_stock=True),
            Shop(id=2, platform_id=2, name="专营店", auto_sync_stock=True),
        ]
    )
    db.add_all(
        [
            Role(id=1, company_id=1, code="WAREHOUSE", name="仓库"),
            Role(id=2, company_id=1, code="ADMIN", name="管理员"),
            Role(id=3, company_id=1, code="QUALITY", name="质检"),
        ]
    )
    db.flush()
    db.add_all(
        [
            StaffRole(staff_id=1, role_id=1),
            StaffRole(staff_id=2, role_id=2),
            StaffRole(staff_id=3, role_id=3),
        ]
    )
    db.commit()


def _staff(db, staff_id: int = 1) -> Staff:
    return db.get(Staff, staff_id)


def _patch_serial(monkeypatch) -> None:
    counter = {"value": 0}

    def fake_serial(prefix, dt=None):
        counter["value"] += 1
        code = prefix.value if hasattr(prefix, "value") else str(prefix)
        return f"{code}20260923-{counter['value']:04d}"

    monkeypatch.setattr(
        stock_settlement_service,
        "generate_serial",
        fake_serial,
    )


def _create_shop_stock(
    db,
    *,
    shop_id: int = 1,
    exclusive_qty: str = "0",
    is_sync: bool = True,
    is_enabled: bool = True,
    effective_start: datetime | None = None,
    effective_end: datetime | None = None,
):
    return stock_settlement_service.create_shop_stock(
        db,
        data={
            "shop_id": shop_id,
            "warehouse_id": 1,
            "product_id": 1,
            "sku_id": 1,
            "exclusive_qty": exclusive_qty,
            "is_sync": is_sync,
            "is_enabled": is_enabled,
            "effective_start": effective_start,
            "effective_end": effective_end,
        },
        staff=_staff(db),
    )


def test_stock_setting_defaults_and_update():
    Session = _session_factory()
    with Session() as db:
        _seed(db)
        row = stock_settlement_service.get_stock_setting(db, 1)
        assert row.sync_stock_default is False
        assert row.shop_stock_enabled is False
        assert row.floating_cost_enabled is False

        updated = stock_settlement_service.update_stock_setting(
            db,
            data={"shop_stock_enabled": True, "floating_cost_enabled": True},
            staff=_staff(db),
        )
        assert updated.shop_stock_enabled is True
        assert updated.floating_cost_enabled is True
        assert updated.sync_stock_default is False

        # 重复读取应命中同一条记录，不重复创建。
        again = stock_settlement_service.get_stock_setting(db, 1)
        assert again.id == updated.id

        with pytest.raises(BizException) as unknown:
            stock_settlement_service.update_stock_setting(
                db,
                data={"unknown_field": True},
                staff=_staff(db),
            )
        assert "不支持的字段" in unknown.value.message


def test_shop_stock_crud_and_exclusive_stock_rule():
    Session = _session_factory()
    with Session() as db:
        _seed(db)
        first = _create_shop_stock(db, shop_id=1, exclusive_qty="4.000")
        data = stock_settlement_service.shop_stock_dict(db, first)
        # 仓库 NORMAL 库存 15，无占用；独享 4 时店铺可用取较小值。
        assert data["stock_qty"] == Decimal("15.000")
        assert data["available_qty"] == Decimal("15.000")
        assert data["shop_available_qty"] == Decimal("4.000")
        assert data["shop_name"] == "旗舰店"
        assert data["shop_type"] == "淘宝"

        # 二号店铺未设置独享额度，不共享仓库剩余库存。
        second = _create_shop_stock(db, shop_id=2, exclusive_qty="0")
        second_data = stock_settlement_service.shop_stock_dict(db, second)
        # 未设置独享额度的店铺不共享仓库剩余库存。
        assert second_data["shop_available_qty"] == Decimal("0.000")
        assert second_data["available_qty"] == Decimal("15.000")

        with pytest.raises(BizException) as conflict:
            _create_shop_stock(db, shop_id=1, exclusive_qty="1.000")
        assert conflict.value.code.value == "40900"

        with pytest.raises(BizException) as invalid_window:
            _create_shop_stock(
                db,
                shop_id=2,
                effective_start=datetime(2026, 9, 2),
                effective_end=datetime(2026, 9, 1),
            )
        assert "有效期开始时间不能晚于结束时间" in invalid_window.value.message

        updated = stock_settlement_service.update_shop_stock(
            db,
            shop_stock_id=first.id,
            data={"exclusive_qty": "8.000"},
            staff=_staff(db),
        )
        assert updated.exclusive_qty == Decimal("8.000")

        listed = stock_settlement_service.list_shop_stocks(db, shop_id=1)
        assert listed["total"] == 1
        assert listed["items"][0]["exclusive_qty"] == Decimal("8.000")

        stock_settlement_service.delete_shop_stock(
            db,
            shop_stock_id=second.id,
            staff=_staff(db),
        )
        assert stock_settlement_service.list_shop_stocks(db)["total"] == 1


def test_sync_shop_stock_requires_setting_and_records_result():
    Session = _session_factory()
    with Session() as db:
        _seed(db)
        _create_shop_stock(db, shop_id=1, exclusive_qty="4.000")
        _create_shop_stock(db, shop_id=2, exclusive_qty="0", is_sync=False)

        with pytest.raises(BizException) as disabled:
            stock_settlement_service.sync_shop_stock(db, staff=_staff(db))
        assert "未启用" in disabled.value.message

        stock_settlement_service.update_stock_setting(
            db,
            data={"shop_stock_enabled": True},
            staff=_staff(db),
        )
        result = stock_settlement_service.sync_shop_stock(db, staff=_staff(db))
        assert result["synced_count"] == 1
        assert result["skipped_count"] == 1
        assert result["synced"][0]["synced_qty"] == Decimal("4.000")
        assert result["skipped"][0]["reason"] == "未开启同步线上库存"

        row = db.scalar(select(ShopStock).where(ShopStock.shop_id == 1))
        assert row is not None
        assert row.last_sync_qty == Decimal("4.000")
        assert row.last_sync_at is not None

        # 未生效（超出有效期）的设置会被跳过。
        stock_settlement_service.update_shop_stock(
            db,
            shop_stock_id=row.id,
            data={"effective_end": datetime(2026, 1, 1)},
            staff=_staff(db),
        )
        expired = stock_settlement_service.sync_shop_stock(db, staff=_staff(db))
        assert expired["synced_count"] == 0
        assert expired["skipped_count"] == 2


def test_settlement_fixes_qty_and_cost_and_rejects_non_month_end():
    Session = _session_factory()
    with Session() as db:
        _seed(db)
        result = stock_settlement_service.create_settlement(
            db,
            staff=_staff(db),
            month="2026-08",
        )
        assert result["settlement_date"].isoformat() == "2026-08-31"
        assert result["item_count"] == 2
        # 结存数量按商品总库存（良品 + 次品）汇总：15 + 2 = 17，配件 3。
        assert result["total_qty"] == Decimal("20.000")
        assert result["total_cost"] == Decimal("185.00")

        rows = {
            row.sku_id: row
            for row in db.scalars(select(StockSettlement))
        }
        assert rows[1].qty == Decimal("17.000")
        assert rows[1].cost_price == Decimal("10.00")
        assert rows[2].qty == Decimal("3.000")
        assert rows[2].cost_price == Decimal("5.00")

        with pytest.raises(BizException) as invalid:
            stock_settlement_service.create_settlement(
                db,
                staff=_staff(db),
                settlement_date="2026-08-15",
            )
        assert "当月最后一天" in invalid.value.message

        # 同一结存日期重复执行按覆盖处理，数量保持稳定。
        again = stock_settlement_service.create_settlement(
            db,
            staff=_staff(db),
            month="2026-08",
        )
        assert again["replaced"] == 2
        assert stock_settlement_service.list_settlements(db)["total"] == 2

        periods = stock_settlement_service.list_settlement_periods(db)
        assert periods[0]["settlement_date"].isoformat() == "2026-08-31"
        assert periods[0]["item_count"] == 2


def test_floating_cost_settlement_uses_period_weighted_average():
    Session = _session_factory()
    with Session() as db:
        _seed(db)
        db.add_all(
            [
                StockFlow(
                    warehouse_id=1,
                    product_id=1,
                    sku_id=1,
                    location_id=1,
                    direction="IN",
                    biz_type="PURCHASE_IN",
                    qty=Decimal("5.000"),
                    cost_price=Decimal("12.00"),
                    total_cost=Decimal("60.00"),
                    flow_time=datetime(2026, 8, 10),
                ),
                StockFlow(
                    warehouse_id=1,
                    product_id=1,
                    sku_id=1,
                    location_id=1,
                    direction="IN",
                    biz_type="PURCHASE_IN",
                    qty=Decimal("5.000"),
                    cost_price=Decimal("16.00"),
                    total_cost=Decimal("80.00"),
                    flow_time=datetime(2026, 8, 20),
                ),
            ]
        )
        db.commit()

        stock_settlement_service.update_stock_setting(
            db,
            data={"floating_cost_enabled": True},
            staff=_staff(db),
        )
        stock_settlement_service.create_settlement(
            db,
            staff=_staff(db),
            month="2026-08",
        )
        row = db.scalar(
            select(StockSettlement).where(StockSettlement.sku_id == 1)
        )
        assert row is not None
        # (60 + 80) / (5 + 5) = 14.00
        assert row.cost_price == Decimal("14.00")
        # 无入库流水的规格仍回退商品规格成本价。
        other = db.scalar(
            select(StockSettlement).where(StockSettlement.sku_id == 2)
        )
        assert other is not None
        assert other.cost_price == Decimal("5.00")


def test_rewrite_outbound_cost_updates_items_and_flows(monkeypatch):
    _patch_serial(monkeypatch)
    Session = _session_factory()
    with Session() as db:
        _seed(db)
        stock_settlement_service.create_settlement(
            db,
            staff=_staff(db),
            month="2026-08",
        )
        stock_out = StockOut(
            id=1,
            stock_out_no="CKD20260923-0001",
            warehouse_id=1,
            stock_out_type="ORDER",
            status="APPROVED",
            audit_status="APPROVED",
            total_qty=Decimal("2.000"),
            created_at=datetime(2026, 9, 10, 9, 0, 0),
        )
        db.add(stock_out)
        db.flush()
        db.add(
            StockOutItem(
                id=1,
                stock_out_id=stock_out.id,
                product_id=1,
                sku_id=1,
                location_id=1,
                qty=Decimal("2.000"),
                cost_price=None,
            )
        )
        db.add_all(
            [
                StockFlow(
                    warehouse_id=1,
                    product_id=1,
                    sku_id=1,
                    location_id=1,
                    direction="OUT",
                    biz_type="ORDER",
                    qty=Decimal("2.000"),
                    flow_time=datetime(2026, 9, 12),
                ),
                StockFlow(
                    warehouse_id=1,
                    product_id=1,
                    sku_id=1,
                    location_id=1,
                    direction="OUT",
                    biz_type="CHECK_OUT",
                    qty=Decimal("1.000"),
                    flow_time=datetime(2026, 9, 13),
                ),
                StockFlow(
                    warehouse_id=1,
                    product_id=1,
                    sku_id=1,
                    location_id=1,
                    direction="IN",
                    biz_type="CHECK_IN",
                    qty=Decimal("1.000"),
                    flow_time=datetime(2026, 9, 14),
                ),
            ]
        )
        db.commit()

        result = stock_settlement_service.rewrite_outbound_cost(
            db,
            staff=_staff(db),
        )
        assert result["period_start"] == datetime(2026, 8, 31)
        assert result["stock_out_item_count"] == 1
        assert result["stock_flow_count"] == 3
        assert result["changed_item_count"] == 1
        assert result["changed_flow_count"] == 3
        assert result["log_count"] == 4
        assert result["batch_no"].startswith("CBHX")

        item = db.get(StockOutItem, 1)
        assert item is not None
        assert item.cost_price == Decimal("10.00")
        flows = list(db.scalars(select(StockFlow)))
        assert all(flow.cost_price == Decimal("10.00") for flow in flows)
        checked_in = next(flow for flow in flows if flow.biz_type == "CHECK_IN")
        assert checked_in.total_cost == Decimal("10.00")

        # 成本回写留痕：记录批次、成本前后值与成本来源。
        logs = stock_settlement_service.list_cost_rewrite_logs(
            db,
            batch_no=result["batch_no"],
        )
        assert logs["total"] == 4
        item_log = next(
            row for row in logs["items"] if row["target_type"] == "STOCK_OUT_ITEM"
        )
        assert item_log["old_cost_price"] is None
        assert item_log["new_cost_price"] == Decimal("10.00")
        assert item_log["new_total_cost"] == Decimal("20.00")
        assert item_log["cost_source"] == "SPEC_COST"
        assert item_log["created_by_name"] == "仓库"
        assert item_log["product_code"] == "P0001"

        # 成本变更后重新回写：从上次结存日到当前时间的出库成本重新计算。
        sku = db.get(ProductSku, 1)
        assert sku is not None
        sku.cost_price = Decimal("12.00")
        db.commit()
        again = stock_settlement_service.rewrite_outbound_cost(
            db,
            staff=_staff(db),
        )
        assert again["changed_item_count"] == 1
        assert again["changed_flow_count"] == 3
        assert again["batch_no"] != result["batch_no"]
        assert db.get(StockOutItem, 1).cost_price == Decimal("12.00")
        latest = stock_settlement_service.list_cost_rewrite_logs(
            db,
            batch_no=again["batch_no"],
        )
        item_log = next(
            row for row in latest["items"] if row["target_type"] == "STOCK_OUT_ITEM"
        )
        assert item_log["old_cost_price"] == Decimal("10.00")
        assert item_log["new_cost_price"] == Decimal("12.00")

        # 已经是目标成本时不再重复计数。
        stable = stock_settlement_service.rewrite_outbound_cost(
            db,
            staff=_staff(db),
        )
        assert stable["changed_item_count"] == 0
        assert stable["changed_flow_count"] == 0
        assert stable["log_count"] == 0


def test_settlement_excel_export_import_and_errors():
    Session = _session_factory()
    with Session() as db:
        _seed(db)
        stock_settlement_service.create_settlement(
            db,
            staff=_staff(db),
            month="2026-08",
        )

        exported = stock_settlement_service.export_settlements(db)
        worksheet = load_workbook(BytesIO(exported)).active
        header = [cell.value for cell in next(worksheet.iter_rows(max_row=1))]
        assert header[0] == "结存日期"
        assert worksheet.max_row == 3

        workbook = load_workbook(
            BytesIO(build_template(stock_settlement_service.SETTLEMENT_TEMPLATE))
        )
        sheet = workbook.active
        sheet.append(
            [
                "2026-07-31",
                "P0001",
                "主商品",
                "P0001-01",
                "默认",
                "7.000",
                "9.00",
            ]
        )
        buffer = BytesIO()
        workbook.save(buffer)
        imported = stock_settlement_service.import_settlements(
            db,
            file_bytes=buffer.getvalue(),
            staff=_staff(db),
        )
        assert imported["created"] == 1
        assert imported["errors"] == []
        row = db.scalar(
            select(StockSettlement).where(
                StockSettlement.settlement_date == datetime(2026, 7, 31)
            )
        )
        assert row is not None
        assert row.qty == Decimal("7.000")
        assert row.cost_price == Decimal("9.00")

        broken = load_workbook(
            BytesIO(build_template(stock_settlement_service.SETTLEMENT_TEMPLATE))
        )
        broken.active.append(
            ["2026-06-30", "NOPE", "不存在", "NOPE-01", "默认", "1.000", "1.00"]
        )
        broken_buffer = BytesIO()
        broken.save(broken_buffer)
        failed = stock_settlement_service.import_settlements(
            db,
            file_bytes=broken_buffer.getvalue(),
            staff=_staff(db),
        )
        assert failed["created"] == 0
        assert failed["errors"]
        assert db.scalar(
            select(StockSettlement).where(
                StockSettlement.settlement_date == datetime(2026, 6, 30)
            )
        ) is None


def _make_api():
    Session = _session_factory(pool=True)

    def override_get_db():
        db = Session()
        try:
            yield db
        finally:
            db.close()

    app = FastAPI()
    register_exception_handlers(app)
    app.include_router(stock_settlement_router)
    app.dependency_overrides[get_db] = override_get_db
    return app, Session


def _auth(staff_id: int, name: str) -> dict:
    return {"Authorization": f"Bearer {create_access_token(staff_id, name, [])}"}


def test_api_stock_settlement_permissions():
    app, Session = _make_api()
    with Session() as db:
        _seed(db)

    client = TestClient(app)
    warehouse_headers = _auth(1, "仓库")
    quality_headers = _auth(3, "质检")

    assert client.get("/stock-settings").status_code == 401

    readable = client.get("/stock-settings", headers=quality_headers)
    assert readable.status_code == 403  # QUALITY 仅虚拟仓，不能读取公司库存配置。
    readable = client.get("/stock-settings", headers=warehouse_headers)
    assert readable.status_code == 200
    assert readable.json()["data"]["shop_stock_enabled"] is False

    denied = client.put(
        "/stock-settings",
        json={"shop_stock_enabled": True},
        headers=quality_headers,
    )
    assert denied.status_code == 403

    enabled = client.put(
        "/stock-settings",
        json={"shop_stock_enabled": True},
        headers=warehouse_headers,
    )
    assert enabled.status_code == 200
    assert enabled.json()["data"]["shop_stock_enabled"] is True

    denied_create = client.post(
        "/shop-stocks",
        json={
            "shop_id": 1,
            "warehouse_id": 1,
            "product_id": 1,
            "sku_id": 1,
            "exclusive_qty": "4.000",
        },
        headers=quality_headers,
    )
    assert denied_create.status_code == 403

    created = client.post(
        "/shop-stocks",
        json={
            "shop_id": 1,
            "warehouse_id": 1,
            "product_id": 1,
            "sku_id": 1,
            "exclusive_qty": "4.000",
            "is_sync": True,
        },
        headers=warehouse_headers,
    )
    assert created.status_code == 200
    payload = created.json()["data"]
    assert payload["shop_available_qty"] == "4.000"
    shop_stock_id = payload["id"]

    synced = client.post(
        "/shop-stocks/sync",
        json={"shop_id": 1},
        headers=warehouse_headers,
    )
    assert synced.status_code == 200
    assert synced.json()["data"]["synced_count"] == 1

    settled = client.post(
        "/stock-settlements",
        json={"month": "2026-08"},
        headers=warehouse_headers,
    )
    assert settled.status_code == 200
    assert settled.json()["data"]["item_count"] == 2

    listed = client.get(
        "/stock-settlements",
        params={"settlement_date": "2026-08-31"},
        headers=quality_headers,
    )
    assert listed.status_code == 200
    assert listed.json()["data"]["total"] == 2

    exported = client.get(
        "/stock-settlements/export",
        headers=quality_headers,
    )
    assert exported.status_code == 200
    assert exported.headers["content-type"].startswith(XLSX_MEDIA_TYPE)

    rewritten = client.post(
        "/stock-settlements/rewrite-outbound-cost",
        json={},
        headers=warehouse_headers,
    )
    assert rewritten.status_code == 200

    logs = client.get(
        "/stock-settlements/cost-rewrite-logs",
        headers=quality_headers,
    )
    assert logs.status_code == 200
    assert logs.json()["data"]["total"] == 0

    deleted = client.delete(
        f"/shop-stocks/{shop_stock_id}",
        headers=warehouse_headers,
    )
    assert deleted.status_code == 200
    assert deleted.json()["message"] == "店铺库存设置已删除"
